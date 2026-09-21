# -*- coding: utf-8 -*-
"""文档解析器：把 PDF / Word 试卷解析成内容块序列，供前端手动分割成题目。

说明：PDF 部分已下沉到 `app.adapters.pdf`（PyMuPDF 的唯一出入口）。
      本模块保留 Word 解析，并对外维持原函数签名，兼容既有脚本。
"""
from docx import Document as DocxDocument

from app.adapters.pdf import parse_blocks as _pdf_parse_blocks


def parse_pdf(path: str) -> list:
    """解析 PDF：按文本块切分，保留页码，图片块标记占位。"""
    return _pdf_parse_blocks(path)


def parse_docx(path: str) -> list:
    """解析 docx：按段落切分，检测段内图片，表格追加在末尾（暂不保持与正文的相对顺序）。"""
    doc = DocxDocument(path)
    blocks = []
    for p in doc.paragraphs:
        xml = p._p.xml
        has_image = ("graphicData" in xml) or ("<pic:pic" in xml)
        if has_image:
            blocks.append({"page": None, "type": "image", "text": "[图片]"})
        text = p.text.strip()
        if text:
            blocks.append({"page": None, "type": "text", "text": text})
    for table in doc.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        blocks.append({"page": None, "type": "table", "text": "[表格]\n" + "\n".join(rows)})
    return blocks
