# -*- coding: utf-8 -*-
"""课表 Excel 导入 + 月历 —— 解析层、匹配/查重、模板、前端守卫。

为什么不起服务
------------------------------------------------
解析是纯函数（services/schedule_import.py 里刻意不 import 任何 app 模块），
匹配/查重/入库是几张表上的纯逻辑（routers/lesson_import.py 里那几个函数），
把函数**直接调用**就能全覆盖 —— 比起 uvicorn 快一个量级，也不占端口、
不会因为「服务没起」而假失败（这个项目在这个坑上栽过）。

分四块：
  ①~⑥ 解析层（日期/时间/时长/金额/枚举/整本工作簿）
  ⑦    模板
  ⑧⑨  匹配、查重、入库（临时库 + 直接调函数，一次 HTTP 都不发）
  ⑩⑪  前端守卫（Vue 真编译器校验模板；日历格子用 node 跑 calendar.js）
  ⑫    一致性守卫（注册顺序、计费唯一出处、依赖声明、文档）

⚠️ 临时库必须在 import app.db **之前**设好 —— engine 绑定在 import 时定死。
   这是本项目自包含测试的标准做法（见 test_batch_page_ai_flow.py）。
"""
from __future__ import annotations

import asyncio
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

TMP = Path(tempfile.gettempdir()) / "shike_schedule_scratch"
shutil.rmtree(TMP, ignore_errors=True)
TMP.mkdir(parents=True, exist_ok=True)
# ⚠️ 必须在 import app.db 之前设
os.environ["SHIKE_DB_PATH"] = str(TMP / "schedule.db")
os.environ["SHIKE_DATA_DIR"] = str(TMP / "data")
os.environ["SHIKE_AUTO_BACKUP"] = "0"

import openpyxl                                     # noqa: E402
from sqlalchemy import select                       # noqa: E402
from starlette.datastructures import UploadFile     # noqa: E402

from app import config                              # noqa: E402
from app.db import SessionLocal, engine             # noqa: E402
from app.models import Base, Lesson, Student        # noqa: E402
from app.routers.lesson_import import (             # noqa: E402
    ImportCommitIn, ImportRowIn, MAX_UPLOAD, commit, preview,
)
from app.services import schedule_import as SI       # noqa: E402

assert str(config.DB_PATH) == str(TMP / "schedule.db"), f"临时库没生效：{config.DB_PATH}"

ROOT = Path(__file__).resolve().parent.parent
FE = ROOT / "frontend"
NODE_MODULES = [
    ROOT / "node_modules" / "jsdom",
    Path("C:/Users/6/.workbuddy/binaries/node/workspace/node_modules/jsdom"),
]

fails: list[str] = []
total = 0


def check(name: str, ok: bool, detail: str = "") -> None:
    global total
    total += 1
    line = f"  [{'OK ' if ok else 'FAIL'}] {name}"
    if detail and not ok:
        line += f"   <- {detail}"
    print(line)
    if not ok:
        fails.append(name)


def section(title: str) -> None:
    print(f"\n==== {title} ====")


def node_bin() -> str | None:
    exe = shutil.which("node")
    return exe


# ---------------------------------------------------------------- 造表工具
def mk_xlsx(rows, sheet: str = "课表", title: str | None = None) -> bytes:
    """把 rows（list[list]）写成 .xlsx 字节，第一行当表头（title 非空则上面再加一行）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    if title:
        ws.append([title])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


HEAD = ["日期", "开始时间", "学生", "时长（分钟）", "本次内容", "上课形式", "地点", "单价", "状态"]


def upload(data: bytes, filename: str = "t.xlsx") -> UploadFile:
    return UploadFile(file=io.BytesIO(data), filename=filename)


# ================================================================ ① 日期
def t_date() -> None:
    section("① 日期怎么写都认")
    f = SI.parse_date
    check("'2026-10-01'", f("2026-10-01") == ("2026-10-01", False), str(f("2026-10-01")))
    check("'2026/10/1' → 补零", f("2026/10/1") == ("2026-10-01", False), str(f("2026/10/1")))
    check("'2026年10月1日'", f("2026年10月1日") == ("2026-10-01", False), str(f("2026年10月1日")))
    check("'2026.10.1'", f("2026.10.1") == ("2026-10-01", False), str(f("2026.10.1")))
    check("Excel 日期单元格（datetime）", f(__import__("datetime").datetime(2026, 10, 1))
          == ("2026-10-01", False))
    check("Excel 序列号 46304 之类能还原", f(46304)[0] is not None, str(f(46304)))
    check("'10-01' 缺年份 → 用默认年并标记", f("10-01", 2026) == ("2026-10-01", True),
          str(f("10-01", 2026)))
    check("'1001' 连写 → 10-01", f("1001", 2026) == ("2026-10-01", True), str(f("1001", 2026)))
    check("'2026-13-01' 非法月 → None", f("2026-13-01")[0] is None, str(f("2026-13-01")))
    check("'2026-02-30' 非法日 → None", f("2026-02-30")[0] is None, str(f("2026-02-30")))
    check("'不是日期' → None（不猜）", f("不是日期")[0] is None, str(f("不是日期")))
    check("空值 → (None, False)", f("") == (None, False) and f(None) == (None, False))


# ================================================================ ② 时间
def t_time() -> None:
    section("② 时间写法 + 区间推时长")
    f = SI.parse_time
    check("'19:00'", f("19:00") == ("19:00", None), str(f("19:00")))
    check("全角「19：00」", f("19：00") == ("19:00", None), str(f("19：00")))
    check("'1900' 连写", f("1900") == ("19:00", None), str(f("1900")))
    check("'930' 连写（补零）", f("930") == ("09:30", None), str(f("930")))
    check("Excel 时间小数 0.7917 ≈ 19:00", f(0.7916666)[0] == "19:00", str(f(0.7916666)))
    check("Excel time 对象", f(__import__("datetime").time(9, 30)) == ("09:30", None))
    check("区间 '19:00-20:30' → 开始 + 90 分钟", f("19:00-20:30") == ("19:00", 90),
          str(f("19:00-20:30")))
    check("区间 '19:00~20:00'（波浪号）", f("19:00~20:00") == ("19:00", 60), str(f("19:00~20:00")))
    check("区间 '19:00 至 20:00'（汉字）", f("19:00 至 20:00") == ("19:00", 60), str(f("19:00 至 20:00")))
    check("跨零点 '21:00-00:30' → 210 分钟", f("21:00-00:30") == ("21:00", 210), str(f("21:00-00:30")))
    check("「周三 19:00」剥掉星期", f("周三 19:00") == ("19:00", None), str(f("周三 19:00")))
    check("「晚上 19:00」剥掉时段词", f("晚上19:00") == ("19:00", None), str(f("晚上19:00")))
    check("'25:99' 非法 → (None, None)", f("25:99") == (None, None), str(f("25:99")))
    check("空值 → (None, None)", f("") == (None, None) and f(None) == (None, None))


# ================================================================ ③ 时长
def t_duration() -> None:
    section("③ 时长怎么写都认")
    f = SI.parse_duration
    check("数字 90", f(90) == (90, None), str(f(90)))
    check("字符串 '90'", f("90") == (90, None), str(f("90")))
    check("'90分钟'", f("90分钟") == (90, None), str(f("90分钟")))
    check("'1.5小时' → 90", f("1.5小时") == (90, None), str(f("1.5小时")))
    check("'1小时30分' → 90", f("1小时30分") == (90, None), str(f("1小时30分")))
    check("'一小时' → 60", f("一小时") == (60, None), str(f("一小时")))
    check("'半小时' → 30", f("半小时") == (30, None), str(f("半小时")))
    check("'两个半小时'... 只认第一个数 → 120（不做通用中文数字解析）",
          f("两个小时") == (120, None), str(f("两个小时")))
    check("空 → (None, None)（由调用方按 60 兜底）", f("") == (None, None) and f(None) == (None, None))
    check("'abc' → 报错文案（不猜）", f("abc")[0] is None and "看不懂" in (f("abc")[1] or ""),
          str(f("abc")))
    check("0 或负 → 报错", f(0)[0] is None and f(-5)[0] is None, f"{f(0)} {f(-5)}")
    check("超过 24 小时 → 报错", f(2000)[0] is None, str(f(2000)))
    check("布尔值 → 报错（不是数字）", f(True)[0] is None, str(f(True)))


# ================================================================ ④ 金额
def t_money() -> None:
    section("④ 单价怎么写都认")
    f = SI.parse_money
    check("数字 300", f(300) == (300.0, None), str(f(300)))
    check("'¥300'", f("¥300") == (300.0, None), str(f("¥300")))
    check("'300元'", f("300元") == (300.0, None), str(f("300元")))
    check("'300/小时' 只取数字", f("300/小时") == (300.0, None), str(f("300/小时")))
    check("'300 元/次' 只取数字", f("300 元/次") == (300.0, None), str(f("300 元/次")))
    check("负数 → 报错", f(-5)[0] is None and "负数" in (f(-5)[1] or ""), str(f(-5)))
    check("'面议' → 报错（不猜）", f("面议")[0] is None, str(f("面议")))
    check("空 → (None, None)（由调用方取学生默认单价）", f("") == (None, None))


# ================================================================ ⑤ 枚举 + 表头
def t_enum() -> None:
    section("⑤ 状态/形式的别名归一 + 表头匹配")
    check("状态「待上」→ scheduled",
          SI.parse_enum("待上", SI.STATUS_ALIASES, "状态")[0] == "scheduled")
    check("状态「已完成」→ done",
          SI.parse_enum("已完成", SI.STATUS_ALIASES, "状态")[0] == "done")
    check("状态「已调课」→ moved",
          SI.parse_enum("已调课", SI.STATUS_ALIASES, "状态")[0] == "moved")
    check("状态「胡思乱想」→ 报错并给出可写项",
          SI.parse_enum("胡思乱想", SI.STATUS_ALIASES, "状态")[0] is None
          and "已排课" in (SI.parse_enum("胡思乱想", SI.STATUS_ALIASES, "状态")[1] or ""))
    check("形式「网课」→ online", SI.parse_enum("网课", SI.MODE_ALIASES, "上课形式")[0] == "online")
    check("形式「面授」→ offline", SI.parse_enum("面授", SI.MODE_ALIASES, "上课形式")[0] == "offline")
    check("必填字段只有三个（日期/开始时间/学生）",
          SI.REQUIRED_FIELDS == ["date", "time", "student"], str(SI.REQUIRED_FIELDS))
    check("列定义里「日期」是必填、「单价」不是",
          [c[2] for c in SI.COLUMNS if c[0] == "date"] == [True]
          and [c[2] for c in SI.COLUMNS if c[0] == "rate"] == [False])

    # 表头匹配：列序调换、括号全半角、加星号都要认
    m = SI._match_columns(["学生", "日期", "开始时间", "时长(分钟)", "备注"])
    check("列序调换照样匹配", m.get("student") == 0 and m.get("date") == 1, str(m))
    check("半角括号「时长(分钟)」认", m.get("duration_min") == 3, str(m))
    m2 = SI._match_columns(["日期*", "开始时间", "学生"])
    check("带 * 的表头认（模板里常这么标必填）",
          {"date", "time", "student"} <= set(m2), str(m2))
    m3 = SI._match_columns(["姓名", "上课日期", "上课时间"])
    check("别名「姓名/上课日期/上课时间」认", {"student", "date", "time"} <= set(m3), str(m3))


# ================================================================ ⑥ 整本工作簿
def t_workbook() -> None:
    section("⑥ 整本工作簿：表头定位 / 缺列 / 忽略列 / 脏行 / 排序")

    r = SI.parse_workbook(b"")
    check("空字节 → ok=False 且有 error", r["ok"] is False and r["error"], str(r.get("error")))
    r = SI.parse_workbook("这不是 Excel".encode("utf-8"))
    check("非 xlsx → ok=False 并提示另存为 .xlsx",
          r["ok"] is False and ".xlsx" in (r["error"] or ""), str(r.get("error")))

    # 表头在第 2 行（上面有标题行）
    data = mk_xlsx([HEAD, ["2026-10-01", "19:00", "小明", 90, "三角函数", "线下", "A", "", "已排课"]],
                   title="10 月课表")
    r = SI.parse_workbook(data)
    check("表头不在第一行也能找到", r["ok"] is True and len(r["rows"]) == 1, str(r.get("error")))
    check("matched_columns 认全 9 列", len(r["matched_columns"]) == 9, str(r["matched_columns"]))

    # 缺必填列 → 整份拒绝
    data = mk_xlsx([["日期", "学生"], ["2026-10-01", "小明"]])
    r = SI.parse_workbook(data)
    check("缺「开始时间」→ ok=False", r["ok"] is False, str(r.get("ok")))
    check("missing_columns 只报缺的那一列", r["missing_columns"] == ["开始时间"],
          str(r["missing_columns"]))
    check("error 里指路「下载模板」", "下载模板" in (r["error"] or ""), str(r.get("error")))

    # 自己加的列被忽略
    data = mk_xlsx([HEAD + ["备注"], ["2026-10-01", "19:00", "小明", "", "", "", "", "", "", "我的笔记"]])
    r = SI.parse_workbook(data)
    check("自己加的「备注」列进 ignored_columns", r["ignored_columns"] == ["备注"],
          str(r["ignored_columns"]))
    check("notes 里说明了忽略了哪些列",
          any("备注" in n for n in r["notes"]), str(r["notes"]))

    # 脏行：错误逐条列出来，不抛异常
    data = mk_xlsx([HEAD, ["不是日期", "25:99", "", "abc", "", "看情况", "", "-5", "胡思乱想"]])
    r = SI.parse_workbook(data)
    check("脏行 → ok=True（文件能读），但该行带 errors",
          r["ok"] is True and len(r["rows"][0]["errors"]) >= 5, str(r["rows"][0]["errors"]))
    errs = " | ".join(r["rows"][0]["errors"])
    for kw in ("日期看不懂", "开始时间看不懂", "没写学生", "看不懂的时长", "看不懂的上课形式",
               "看不懂的状态", "负数"):
        check(f"脏行错误里点出「{kw}」", kw in errs, errs)

    # 空行跳过 + 空表
    data = mk_xlsx([HEAD, [], ["2026-10-01", "19:00", "小明"]])
    r = SI.parse_workbook(data)
    check("整行空被跳过（不是一条错行）", len(r["rows"]) == 1, str(len(r["rows"])))
    data = mk_xlsx([HEAD])
    r = SI.parse_workbook(data)
    check("只有表头 → ok=False 且提示没有数据行",
          r["ok"] is False and "没有可读的数据行" in (r["error"] or ""), str(r.get("error")))

    # 排序：能导的在前、有错的沉底
    data = mk_xlsx([HEAD,
                    ["2026-10-05", "10:00", "小明"],
                    ["坏", "", "小明"],
                    ["2026-10-02", "09:00", "小明"]])
    r = SI.parse_workbook(data)
    check("能导的行按时间排在前面",
          [bool(x["errors"]) for x in r["rows"]] == [False, False, True],
          str([(x["row_no"], bool(x["errors"])) for x in r["rows"]]))
    check("有错的行是最后一条且保留了原行号", r["rows"][-1]["row_no"] == 3,
          str(r["rows"][-1]["row_no"]))

    # 缺年份的提示
    data = mk_xlsx([HEAD, ["10-01", "19:00", "小明"]])
    r = SI.parse_workbook(data, default_year=2026)
    check("补年份的行 date=2026-10-01", r["rows"][0]["date"] == "2026-10-01", str(r["rows"][0]))
    check("notes 里说明按哪一年处理",
          any("2026 年" in n for n in r["notes"]), str(r["notes"]))

    # 时长从区间推
    data = mk_xlsx([HEAD, ["2026-10-01", "19:00-20:30", "小明"]])
    r = SI.parse_workbook(data)
    check("区间推出来的时长 90 落到 duration_min", r["rows"][0]["duration_min"] == 90,
          str(r["rows"][0]))
    check("notes 里说明时长是算出来的",
          any("区间" in n for n in r["notes"]), str(r["notes"]))

    # 多工作表：优先选必填列认全的那个
    data = mk_xlsx([["甲", "乙"], ["1", "2"]], sheet="杂表")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wb.create_sheet("课表")
    ws = wb["课表"]
    for r_ in [HEAD, ["2026-10-01", "19:00", "小明"]]:
        ws.append(r_)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    r = SI.parse_workbook(buf.getvalue())
    check("多个工作表时挑出表头对的那个", r["sheet"] == "课表", str(r["sheet"]))
    check("sheets 里列出全部工作表名", r["sheets"] == ["杂表", "课表"], str(r["sheets"]))

    # 大文件上限
    check("MAX_ROWS 是个合理上限（几百到几千）", 100 <= SI.MAX_ROWS <= 10000, str(SI.MAX_ROWS))

    # parse_workbook 永不抛异常
    for junk in (b"PK\x03\x04garbage", b"\x00\x01\x02", b"<html></html>"):
        try:
            rr = SI.parse_workbook(junk)
            ok = rr["ok"] is False
        except Exception as e:                       # pragma: no cover
            ok = False
            rr = {"error": repr(e)}
        check(f"垃圾输入不抛异常（{junk[:8]!r}）", ok, str(rr.get("error")))


# ================================================================ ⑦ 模板
def t_template() -> None:
    section("⑦ 导入模板")
    blob = SI.build_template()
    check("返回的是 xlsx（ZIP 头 PK）", blob[:2] == b"PK", blob[:8])
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    check("两个工作表：课表 + 填写说明", wb.sheetnames == ["课表", "填写说明"], str(wb.sheetnames))
    ws = wb["课表"]
    hdr = [c.value for c in ws[1]]
    check("表头就是 COLUMNS 的显示名（唯一出处，不会文档说 A 代码认 B）",
          hdr == SI.TEMPLATE_HEADERS, str(hdr))
    check("表头 9 列且含三个必填列",
          len(hdr) == 9 and all(k in hdr for k in ("日期", "开始时间", "学生")), str(hdr))
    check("冻结首行", ws.freeze_panes == "A2", str(ws.freeze_panes))
    check("含示例行（第 2、3 行）", ws.cell(row=2, column=1).value is not None)
    check("示例行下面写了「填之前请删掉」",
          any("示例" in str(ws.cell(row=r, column=1).value or "") for r in range(1, 8)),
          str([ws.cell(row=r, column=1).value for r in range(1, 8)]))

    ws2 = wb["填写说明"]
    col_a = [ws2.cell(row=r, column=1).value for r in range(1, ws2.max_row + 1)]
    col_b = [ws2.cell(row=r, column=2).value for r in range(1, ws2.max_row + 1)]
    check("说明页有小标题「这份表怎么填」", "这份表怎么填" in col_a, str(col_a[:5]))
    check("说明页有小标题「每一列」", "每一列" in col_a, str(col_a[:10]))
    check("说明页有小标题「几种写法都认」", "几种写法都认" in col_a)
    check("说明页有小标题「导入时会发生什么」", "导入时会发生什么" in col_a)
    check("每个字段都在说明页出现（含必填/可空标注）",
          all(any(f"「{label}」" in str(v) for v in col_a)
              for _f, label, _r, _a, _h in SI.COLUMNS),
          str([v for v in col_a if v and "「" in str(v)])[:4])
    check("A 列没有被塞进长段说明（列序没颠倒）",
          not any(v and len(str(v)) > 60 for v in col_a),
          str([v for v in col_a if v and len(str(v)) > 60][:2]))
    check("长段说明都在 B 列", sum(1 for v in col_b if v and len(str(v)) > 20) >= 6,
          str(sum(1 for v in col_b if v and len(str(v)) > 20)))
    check("说明里写到了 .xls 要另存为 .xlsx（或至少提到格式）",
          any(v and "xlsx" in str(v) for v in col_b), "")
    wb.close()

    # 模板自己必须能被自己解析回来（防止「生成的模板导不进去」）
    r = SI.parse_workbook(blob)
    check("生成的模板能被 parse_workbook 读回（表头对得上）",
          r["ok"] is True and r["missing_columns"] == [], str(r.get("error")))
    check("模板里的示例行标记成「示例：…」的学生（不会被当成真学生）",
          any("示例" in (x["student"] or "") for x in r["rows"]),
          str([x["student"] for x in r["rows"]]))


# ================================================================ ⑧⑨ 路由层
def seed(db) -> tuple[int, int]:
    """建两个学生 + 一条已排的课。返回 (有默认单价的学生 id, 没单价的)。"""
    a = Student(owner_id=config.OWNER_ID, name="小明", hourly_rate=300.0,
                rate_unit="hour", status="active")
    b = Student(owner_id=config.OWNER_ID, name="小红", hourly_rate=None,
                rate_unit="hour", status="active")
    db.add_all([a, b])
    db.flush()
    db.add(Lesson(owner_id=config.OWNER_ID, student_id=a.id, start_at="2026-10-05T19:00",
                  duration_min=60, status="done", mode="offline", rate=300.0,
                  billable=1, amount=300.0, node_ids="[]"))
    db.commit()
    return a.id, b.id


def t_preview(sid_a: int, sid_b: int) -> dict:
    section("⑧ 预览：匹配 / 查重（含表内重复）/ 单价 / 计数")
    rows = [
        HEAD,
        ["2026-10-05", "19:00", "小明", 60, "和库里那条撞了"],          # 行2 → 与库重复
        ["2026-10-05", "19:00", "小明", 60, "表里也写了一遍"],          # 行3 → 表内重复（指回行2）
        ["2026-10-06", "10:00", "小红", "", "小红没默认单价"],          # 行4 → 能导但无单价
        ["2026-10-07", "10:00", "小明", 90, ""],                        # 行5 → 能导，单价 300
        ["2026-10-08", "10:00", "查无此人", 60, ""],                    # 行6 → 学生对不上
        ["坏", "", "", "", ""],                                          # 行7 → 有错
    ]
    data = mk_xlsx(rows)
    db = SessionLocal()
    try:
        res = asyncio.run(preview(file=upload(data), db=db))
    finally:
        db.close()

    check("ok=True", res.get("ok") is True, str(res.get("error")))
    c = res["counts"]
    check("total = 6", c["total"] == 6, str(c))
    check("ready = 2（行4、行5）", c["ready"] == 2, str(c))
    check("duplicate = 2", c["duplicate"] == 2, str(c))
    check("duplicate_in_file = 1（只有行3 是表里写重）", c["duplicate_in_file"] == 1, str(c))
    check("no_student = 1", c["no_student"] == 1, str(c))
    check("problem = 1", c["problem"] == 1, str(c))
    check("no_rate = 2（小红那行 + 学生不存在那行 —— 后者建了学生也没默认单价）",
          c["no_rate"] == 2, str(c))
    check("unknown_students 列出「查无此人」", res["unknown_students"] == ["查无此人"],
          str(res["unknown_students"]))
    check("months = ['2026-10']", res["months"] == ["2026-10"], str(res["months"]))

    by = {r["row_no"]: r for r in res["rows"]}
    check("行2 与库里那条重复 → duplicate_kind=db",
          by[2]["action"] == "duplicate" and by[2]["duplicate_kind"] == "db",
          f"{by[2]['action']} / {by[2].get('duplicate_kind')}")
    check("行2 duplicate_of 给的是库里的课 id",
          isinstance((by[2]["duplicate_of"] or {}).get("id"), int), str(by[2].get("duplicate_of")))
    check("行3 表内重复 → duplicate_kind=file",
          by[3]["action"] == "duplicate" and by[3]["duplicate_kind"] == "file",
          f"{by[3]['action']} / {by[3].get('duplicate_kind')}")
    check("行3 duplicate_of.in_file_row = 2（指回第一次出现的那行）",
          (by[3]["duplicate_of"] or {}).get("in_file_row") == 2, str(by[3].get("duplicate_of")))
    check("行4 学生存在（小红）→ action=import",
          by[4]["action"] == "import" and by[4]["student_exists"] is True, str(by[4]["action"]))
    check("行4 rate_effective = None（学生档案没默认单价、表里也没写）",
          by[4]["rate_effective"] is None, str(by[4].get("rate_effective")))
    check("行5 rate_effective = 300（学生档案里的默认单价）",
          by[5]["rate_effective"] == 300.0, str(by[5].get("rate_effective")))
    check("行6 action=no_student", by[6]["action"] == "no_student", str(by[6]["action"]))
    check("行7 action=blocked 且不被勾选", by[7]["action"] == "blocked", str(by[7]["action"]))
    check("每行都有 duplicate_kind / rate_effective / student_id / action 键",
          all(k in r for r in res["rows"]
              for k in ("duplicate_kind", "rate_effective", "student_id", "action")))
    notes = " | ".join(res["notes"])
    check("notes 说明「在表里就写重了」", "写重" in notes, notes)
    check("notes 说明「和已排的课撞了」", "撞了" in notes, notes)
    check("notes 说明「定不出单价」", "定不出单价" in notes, notes)
    check("notes 说明「这些学生系统里没有」", "系统里没有" in notes, notes)

    # 缺列 / 非 Excel 也走同一个函数，返回 ok=False 而不是抛
    bad = asyncio.run(preview(file=upload(mk_xlsx([["日期", "学生"], ["2026-10-01", "小明"]])),
                              db=SessionLocal()))
    check("缺列时返回 ok=False（不抛异常）", bad["ok"] is False, str(bad.get("ok")))
    return res


def t_commit(sid_a: int, sid_b: int, res: dict) -> None:
    section("⑨ 入库：写库 / 跳过 / 顺手建学生 / 单价快照")
    db = SessionLocal()
    try:
        n_lesson0 = len(db.scalars(select(Lesson)).all())
        n_stud0 = len(db.scalars(select(Student)).all())

        picked = [r for r in res["rows"] if r["action"] in ("import", "no_student")]
        check("挑出 3 行来提交（2 可导 + 1 学生对不上）", len(picked) == 3, str(len(picked)))
        payload = ImportCommitIn(rows=[
            ImportRowIn(row_no=r["row_no"], date=r["date"], time=r["time"], student=r["student"],
                        student_id=r["student_id"], duration_min=r["duration_min"],
                        topic=r["topic"], mode=r["mode"], location=r["location"],
                        rate=r["rate"], status=r["status"])
            for r in picked
        ], create_students=False)
        rep = commit(payload, db)
        check("created = 2", rep["created"] == 2, str(rep))
        check("skipped = 1 且原因是学生不存在",
              rep["skipped"] == 1 and "学生" in rep["skipped_rows"][0]["reason"], str(rep["skipped_rows"]))
        check("没勾「建学生」→ created_students 为空", rep["created_students"] == [], str(rep))
        check("库里真的多了 2 条",
              len(db.scalars(select(Lesson)).all()) == n_lesson0 + 2,
              f"{n_lesson0} -> {len(db.scalars(select(Lesson)).all())}")
        check("学生数没变", len(db.scalars(select(Student)).all()) == n_stud0, "")

        made = db.scalars(select(Lesson).where(Lesson.start_at == "2026-10-07T10:00")).all()
        check("入库的那节课在库里找得到", len(made) == 1, str(len(made)))
        check("90 分钟 × 300 元/小时 → amount=450（复用 lessons._apply_billing）",
              made[0].amount == 450.0, str(made[0].amount))
        check("单价快照：rate 记为 300", made[0].rate == 300.0, str(made[0].rate))
        check("billable = 1", made[0].billable == 1, str(made[0].billable))
        check("node_ids 落成 '[]'（不是 NULL，前端直接 JSON.parse）",
              made[0].node_ids == "[]", str(made[0].node_ids))

        low = db.scalars(select(Lesson).where(Lesson.start_at == "2026-10-06T10:00")).all()
        check("小红那节按时长 60 兜底（表里没写时长）", low[0].duration_min == 60, str(low[0].duration_min))
        check("小红没有默认单价 → amount 为空（与手工排课一致）",
              low[0].amount is None and low[0].rate is None, f"{low[0].rate} {low[0].amount}")

        # 顺手建学生
        before = len(db.scalars(select(Student)).all())
        rep2 = commit(ImportCommitIn(rows=[ImportRowIn(
            row_no=99, date="2026-10-20", time="15:00", student="新学生甲",
            student_id=None, duration_min=60, topic="第一次课", mode="offline",
            location="", rate=None, status="scheduled")], create_students=True), db)
        check("勾了建学生 → created = 1", rep2["created"] == 1, str(rep2))
        check("created_students 回报了名字", rep2["created_students"] == ["新学生甲"], str(rep2))
        check("学生表真的多了一个",
              len(db.scalars(select(Student)).all()) == before + 1, "")
        ns = db.scalar(select(Student).where(Student.name == "新学生甲"))
        check("新学生 status=active", ns is not None and ns.status == "active",
              str((ns.name, ns.status) if ns else None))

        # 同一批里重名的学生只建一个
        before2 = len(db.scalars(select(Student)).all())
        rep3 = commit(ImportCommitIn(rows=[
            ImportRowIn(row_no=100, date="2026-10-21", time="15:00", student="新学生乙",
                        duration_min=60),
            ImportRowIn(row_no=101, date="2026-10-22", time="15:00", student="新学生乙",
                        duration_min=60),
        ], create_students=True), db)
        check("同一批里重名的学生只建一个（created=2 但学生只 +1）",
              rep3["created"] == 2 and len(db.scalars(select(Student)).all()) == before2 + 1,
              str(rep3["created_students"]))

        # 非法输入：宁可跳过，不写脏数据
        rep4 = commit(ImportCommitIn(rows=[ImportRowIn(
            row_no=110, date="2026-13-45", time="99:99", student="小明",
            student_id=sid_a)]), db)
        check("日期时间非法 → 跳过而不是写脏数据",
              rep4["created"] == 0 and rep4["skipped"] == 1, str(rep4))
        rep5 = commit(ImportCommitIn(rows=[ImportRowIn(
            row_no=111, date="2026-10-30", time="09:00", student="小明", student_id=999999)]), db)
        check("学生 id 不存在 → 跳过", rep5["created"] == 0 and rep5["skipped"] == 1, str(rep5))
        rep6 = commit(ImportCommitIn(rows=[ImportRowIn(
            row_no=112, date="2026-10-30", time="09:00", student="小明",
            student_id=sid_a, status="胡来", mode="胡来")]), db)
        check("状态/形式写了怪值 → 回落成 scheduled/offline 而不是写进库里",
              rep6["created"] == 1
              and db.scalar(select(Lesson).where(Lesson.start_at == "2026-10-30T09:00")).status
              == "scheduled", str(rep6))

        try:
            commit(ImportCommitIn(rows=[]), db)
            check("空 rows → 抛 422", False, "没抛")
        except Exception as e:
            check("空 rows → 抛 422", getattr(e, "status_code", None) == 422, repr(e))
    finally:
        db.close()


# ================================================================ ⑩ 前端守卫
def extract_template(path: Path) -> str | None:
    src = path.read_text(encoding="utf-8")
    key = "template: `"
    i = src.find(key)
    if i < 0:
        return None
    j = src.find("`", i + len(key))
    if j < 0:
        return None
    return src[i + len(key):j]


def t_frontend() -> None:
    section("⑩ 前端守卫（Vue 真编译器校验模板 + 关键结构）")
    cal = FE / "src" / "calendar.js"
    sched = FE / "src" / "views" / "Schedule.js"
    dlg = FE / "src" / "components" / "ScheduleImportDialog.js"
    api = FE / "src" / "api.js"
    css = FE / "app.css"
    for p in (cal, sched, dlg, api, css):
        check(f"{p.relative_to(ROOT)} 存在", p.exists(), str(p))

    src = cal.read_text(encoding="utf-8")
    for fn in ("parsePeriod", "daysInMonth", "monthGrid", "shiftPeriod", "periodLabel",
               "lessonsByDate", "monthDates", "WEEK_HEADS"):
        check(f"calendar.js 导出 {fn}", re.search(rf"export\s+(const|function)\s+{fn}\b", src) is not None)
    check("calendar.js 不碰 DOM / 不 import 任何东西（纯逻辑才好单测）",
          "document" not in src and not re.search(r"^\s*import\s", src, re.M),
          "")

    src = sched.read_text(encoding="utf-8")
    check("Schedule.js 从 ../calendar.js 引入月历函数",
          "from '../calendar.js'" in src, "")
    check("Schedule.js 引入了 ScheduleImportDialog",
          "ScheduleImportDialog" in src, "")
    check("有「月历 / 列表」两个视图切换按钮",
          ">月历</button>" in src and ">列表</button>" in src,
          str(re.findall(r">([月列][历表])</button>", src)))
    check("默认视图是月历", re.search(r"view\s*=\s*ref\(\s*'calendar'\s*\)", src) is not None,
          str(re.search(r"view\s*=\s*ref\([^)]*\)", src)))
    check("月历里用 grid 渲染格子", 'v-for="(c, i) in grid"' in src, "")
    check("点空格能排课（openCreate 带日期）", "openCreate(c.date)" in src, "")
    check("翻月用 shiftPeriod", "shiftPeriod(" in src, "")

    src = dlg.read_text(encoding="utf-8")
    check("导入弹窗三步：pick/preview/done",
          all(f"step === '{s}'" in src for s in ("pick", "preview", "done")), "")
    check("有「下载导入模板」入口", "下载导入模板" in src, "")
    check("有「顺手把这几个学生还建到系统里」开关", "顺手把这几个学生" in src, "")
    check("默认只勾 action=import 的行", "r.action === 'import'" in src, "")
    check("表内重复与库重复文案不同（dupText）", "dupText" in src and "表内写重了" in src, "")
    check("无单价的行有提示", "没有单价" in src, "")

    src = api.read_text(encoding="utf-8")
    check("api.js 有 scheduleImportApi.templateUrl/preview/commit",
          all(k in src for k in ("scheduleImportApi", "templateUrl:", "preview:", "commit:")), "")
    check("preview 走 multipart（api.upload，不能手写 JSON）",
          re.search(r"preview:\s*\(formData\)\s*=>\s*api\.upload\(", src) is not None, "")

    src = css.read_text(encoding="utf-8")
    for cls in (".cal-nav", ".cal-head", ".cal-grid", ".cal-cell", ".cal-item", ".cal-add"):
        # 允许它出现在「.cal-head, .cal-grid { ... }」这种并列选择器里，不要求在行首
        check(f"app.css 有 {cls}",
              re.search(rf"(?:^|[,{{]\s*){re.escape(cls)}\b", src, re.M) is not None, "")
    check("填位格 / 今天 有区分样式（.cal-cell.out / .cal-cell.today）",
          ".cal-cell.out" in src and ".cal-cell.today" in src, "")

    # ---- Vue 真编译器：抓 compiler-30（孤立的 v-else-if）这类模板错误 ----
    # ⚠️ 这里有两个坑，都踩过：
    #   ① Vue 的 compiler-dom 在遇到属性值里的 `&`（比如 `a && b.length`）时会去
    #      decodeEntities → 需要一个 document。Node 里没有 → 报
    #      `ReferenceError: document is not defined`，看起来像模板有问题，其实是环境问题。
    #      所以下面给一个只够它用的 document stub。
    #   ② 正因为它可能因为环境问题而**假通过/假失败**，脚本里额外编译一份
    #      **故意写错**的模板（孤立的 v-else-if），必须报失败 —— 用这条证明
    #      「编译器真的在跑」，而不是被 stub 兜成永远 OK。
    node = node_bin()
    if not node:
        print("  [SKIP] node 不在 PATH 上，跳过模板编译校验")
        return
    vue = FE / "vendor" / "vue.global.prod.js"
    check("vendor/vue.global.prod.js 存在", vue.exists(), str(vue))
    if not vue.exists():
        return

    js = TMP / "compile_check.js"
    js.write_text(
        "const fs=require('fs'),vm=require('vm');\n"
        f"const VUE={str(vue)!r};\n"
        "function mkDoc() {\n"
        "  const el=()=>{const o={_raw:'',_foo:'',_text:''};\n"
        "    Object.defineProperty(o,'innerHTML',{set(v){o._raw=v;const m=/foo=\"([\\s\\S]*)\"/.exec(v);"
        "o._foo=m?m[1]:'';o._text=v;},get(){return o._raw;}});\n"
        "    Object.defineProperty(o,'children',{get(){const dec=s=>String(s).replace(/&quot;/g,'\"')"
        ".replace(/&#39;/g,\"'\").replace(/&lt;/g,'<').replace(/&gt;/g,'>').replace(/&amp;/g,'&');\n"
        "      return [{getAttribute:()=>dec(o._foo)}];}});\n"
        "    Object.defineProperty(o,'textContent',{get(){return String(o._text)"
        ".replace(/&lt;/g,'<').replace(/&gt;/g,'>').replace(/&amp;/g,'&').replace(/&quot;/g,'\"');}});\n"
        "    return o;};\n"
        "  return {createElement:()=>el(),querySelector:()=>null,createTextNode:()=>({})};\n"
        "}\n"
        "const ctx={console,document:mkDoc(),window:{}};vm.createContext(ctx);\n"
        "vm.runInContext(fs.readFileSync(VUE,'utf8'),ctx);\n"
        "function comp(tpl){try{ctx.Vue.compile(tpl);return 'OK';}"
        "catch(e){return 'FAIL '+String(e.message).slice(0,120);}}\n"
        "// 自检：这份模板孤立 v-else-if，必须报错，否则说明校验是空转\n"
        "console.log('SELFTEST ' + comp('<div><template v-else-if=\"b\">x</template></div>'));\n"
        "console.log('SELFTEST_GOOD ' + comp('<div><p v-if=\"a\">1</p><p v-else>2</p></div>'));\n"
        "for (const f of process.argv.slice(2)) console.log('TPL ' + f + ' ' + comp(fs.readFileSync(f,'utf8')));\n",
        encoding="utf-8")

    tpl_files = []
    for path in (sched, dlg):
        tpl = extract_template(path)
        check(f"{path.name} 能提取出 template（模板写成字符串常量）",
              bool(tpl) and tpl.rstrip().endswith(">"), (tpl or "")[-60:])
        if not tpl:
            continue
        check(f"{path.name} 的模板里没有 <script>（in-dom 模板不带脚本）", "<script" not in tpl, "")
        tf = TMP / f"tpl_{path.stem}.html"
        tf.write_text(tpl, encoding="utf-8")
        tpl_files.append((path, tf))

    r = subprocess.run([node, str(js)] + [str(t) for _p, t in tpl_files],
                       capture_output=True, text=True, timeout=90)
    out = (r.stdout or "") + (r.stderr or "")
    check("编译器自检：故意写错的模板（孤立 v-else-if）必须报错",
          any(l.startswith("SELFTEST ") and l.startswith("SELFTEST FAIL") for l in out.splitlines()),
          out.strip()[:200])
    check("编译器自检：正常模板必须通过",
          "SELFTEST_GOOD OK" in out, out.strip()[:200])
    for path, tf in tpl_files:
        line = [l for l in out.splitlines() if l.startswith(f"TPL {tf} ")]
        got = line[0].split(" ", 2)[2] if line else "（没有输出）"
        check(f"{path.name} 的模板能被 Vue 编译器吃下（抓 compiler-30 这类错）",
              got == "OK", got)


# ================================================================ ⑪ 日历格子
CAL_JS = r"""
import { WEEK_HEADS, parsePeriod, daysInMonth, monthGrid, shiftPeriod, periodLabel,
         lessonsByDate, monthDates } from './calendar.mjs';
const R = [];
const eq = (n, a, b) => R.push([n, JSON.stringify(a) === JSON.stringify(b), JSON.stringify(a) + ' vs ' + JSON.stringify(b)]);

eq('周首是周日开头的七列', WEEK_HEADS.length, 7);
eq('第一列是周日', WEEK_HEADS[0], '周日');
eq('最后一列是周六', WEEK_HEADS[6], '周六');

eq('parsePeriod 正常', parsePeriod('2026-10'), { year: 2026, month: 10 });
eq('parsePeriod 单位数月份', parsePeriod('2026-1'), { year: 2026, month: 1 });
eq('parsePeriod 13 月非法', parsePeriod('2026-13'), null);
eq('parsePeriod 0 月非法', parsePeriod('2026-0'), null);
eq('parsePeriod 垃圾', parsePeriod('abc'), null);
eq('parsePeriod 空', parsePeriod(''), null);
eq('parsePeriod 去空格', parsePeriod(' 2026-10 '), { year: 2026, month: 10 });

eq('10 月 31 天', daysInMonth('2026-10'), 31);
eq('2 月平年 28 天', daysInMonth('2026-02'), 28);
eq('2 月闰年 29 天', daysInMonth('2024-02'), 29);
eq('4 月 30 天', daysInMonth('2026-04'), 30);

const g = monthGrid('2026-10');
eq('2026-10 网格长度 35（5 周）', g.length, 35);
eq('2026-10-01 是周四 → 前面填 4 个空位', g.filter(c => c.out).length, 35 - 31);
eq('前 4 格是填位', g.slice(0, 4).every(c => c.date === null && c.out === true), true);
eq('第 5 格是 10-01', g[4].date, '2026-10-01');
eq('第 5 格的 day 是 1', g[4].day, 1);
eq('10-01 不是填位格', g[4].out, false);
eq('10-31 在网格里', g.some(c => c.date === '2026-10-31'), true);
eq('网格长度是 7 的倍数（每种月份都要成立）',
   ['2026-01','2026-02','2026-04','2026-10','2026-11','2026-12','2024-02','2023-02']
     .every(p => monthGrid(p).length % 7 === 0), true);
eq('2026-02-01 是周日 → 没有前导空位',
   monthGrid('2026-02').filter(c => c.out && c.date === null).length, 0);
eq('2026-11 也正好铺满（1 号周日 + 30 天）', monthGrid('2026-11').length, 35);
eq('非法 period → 空网格', monthGrid('nope').length, 0);

eq('往后翻一个月', shiftPeriod('2026-10', 1), '2026-11');
eq('往前翻一个月', shiftPeriod('2026-10', -1), '2026-09');
eq('12 月往后跨年', shiftPeriod('2026-12', 1), '2027-01');
eq('1 月往前跨年', shiftPeriod('2026-01', -1), '2025-12');
eq('非法 period 原样返回', shiftPeriod('nope', 1), 'nope');
eq('label 中文', periodLabel('2026-10'), '2026 年 10 月');
eq('label 非法原样', periodLabel('nope'), 'nope');

const ls = [
  { id: 1, start_at: '2026-10-01T19:00' },
  { id: 2, start_at: '2026-10-01T09:00' },
  { id: 3, start_at: '2026-10-03T10:00' },
  { id: 4, start_at: '垃圾数据' },
  { id: 5 },
];
const map = lessonsByDate(ls);
eq('丢掉没有合法日期的行', map.size, 2);
eq('同一天按开始时间排序', map.get('2026-10-01').map(x => x.id), [2, 1]);
eq('另一天只有一个', map.get('2026-10-03').length, 1);
eq('空输入不炸', lessonsByDate(null).size, 0);
eq('monthDates 长度 = 当月天数', monthDates('2026-10').length, 31);
eq('monthDates 首尾正确',
   [monthDates('2026-10')[0], monthDates('2026-10')[30]], ['2026-10-01', '2026-10-31']);

let bad = 0;
for (const [n, ok, d] of R) {
  console.log('  [' + (ok ? 'OK ' : 'FAIL') + '] ' + n + (ok ? '' : '   <- ' + d));
  if (!ok) bad += 1;
}
console.log('__CAL_DONE__ ' + R.length + ' ' + bad);
"""


def t_calendar() -> None:
    section("⑪ 日历格子（node 直接跑 calendar.js）")
    node = node_bin()
    if not node:
        print("  [SKIP] node 不在 PATH 上")
        return
    src = FE / "src" / "calendar.js"
    check("calendar.js 存在（不然日历无从测起）", src.exists(), str(src))
    if not src.exists():
        return
    d = TMP / "cal"
    d.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, d / "calendar.mjs")
    (d / "run.mjs").write_text(CAL_JS, encoding="utf-8")
    r = subprocess.run([node, str(d / "run.mjs")], capture_output=True, text=True, timeout=60)
    out = (r.stdout or "") + (r.stderr or "")
    for line in out.splitlines():
        if line.startswith("  [") and ("[OK ]" in line or "[FAIL]" in line):
            m = re.match(r"\s*\[(OK |FAIL)\]\s*(.+?)(?:\s+<- (.*))?$", line)
            if m:
                check(m.group(2).strip(), m.group(1) == "OK ", m.group(3) or "")
    tail = [l for l in out.splitlines() if l.startswith("__CAL_DONE__")]
    check("node 侧日历断言全部跑完（进程没中途挂掉）", len(tail) == 1, out.strip()[-200:])
    if tail:
        n, bad = tail[0].split()[1:3]
        check(f"node 侧 {n} 项里 0 失败", bad == "0", f"失败 {bad} 项")


# ================================================================ ⑫ 一致性守卫
def t_consistency() -> None:
    section("⑫ 一致性守卫（注册顺序 / 计费唯一出处 / 依赖 / 文档）")
    main_src = (ROOT / "backend" / "main.py").read_text(encoding="utf-8")
    i_imp = main_src.find("lesson_import.router)")
    i_les = main_src.find("lessons.router)")
    check("main.py 注册了 lesson_import.router", i_imp > 0, "")
    check("lesson_import 注册在 lessons **之前**（否则 /import 可能被 /{lid} 吃掉）",
          0 < i_imp < i_les, f"{i_imp} vs {i_les}")

    li = (ROOT / "backend" / "app" / "routers" / "lesson_import.py").read_text(encoding="utf-8")
    check("入库复用 lessons._apply_billing（不自己算钱，免得两处算法漂移）",
          "from .lessons import _apply_billing" in li and "_apply_billing(db, ls)" in li, "")
    check("没有自己算 amount（出现 '* 60' 之类就是漂移了）",
          not re.search(r"amount\s*=", li), "")
    check("有上传大小上限", "MAX_UPLOAD" in li, "")
    check("状态/形式在入库前再校验一次（前端可以被改）",
          "ALLOWED_STATUS" in li and "ALLOWED_MODE" in li, "")

    si = (ROOT / "backend" / "app" / "services" / "schedule_import.py").read_text(encoding="utf-8")
    check("解析层不 import 任何 app.* 模块（保持纯逻辑，才能脱离服务单测）",
          not re.search(r"^\s*from\s+\.+", si, re.M), "")
    check("openpyxl 缺失时不让服务起不来（try/except ImportError）",
          "_HAS_OPENPYXL" in si and "ImportError" in si, "")
    check("列定义只有一处（COLUMNS），模板与文档都从它取",
          si.count("COLUMNS: list[") == 1 and "TEMPLATE_HEADERS = [label for _f, label" in si, "")

    req = (ROOT / "backend" / "requirements.txt").read_text(encoding="utf-8")
    check("requirements.txt 声明了 openpyxl", "openpyxl" in req, "")

    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    check("README 有课表导入的 Excel 格式说明",
          "课表" in readme and ("Excel" in readme or "xlsx" in readme), "")
    check("README 说明了那三列是必填",
          "日期" in readme and "开始时间" in readme and "学生" in readme, "")


# ================================================================ main
def main() -> int:
    Base.metadata.create_all(engine)          # 临时库从零建表（迁移链另有测试守着）
    db = SessionLocal()
    try:
        sid_a, sid_b = seed(db)
    finally:
        db.close()

    t_date()
    t_time()
    t_duration()
    t_money()
    t_enum()
    t_workbook()
    t_template()
    res = t_preview(sid_a, sid_b)
    t_commit(sid_a, sid_b, res)
    t_frontend()
    t_calendar()
    t_consistency()

    print("\n" + "=" * 64)
    print(f"  共 {total} 项，失败 {len(fails)} 项")
    if fails:
        for n in fails:
            print(f"    ✗ {n}")
    print("=" * 64)
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
