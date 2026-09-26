# -*- coding: utf-8 -*-
"""笔记导出：Markdown -> Word(.docx) / PDF。

和 paper_export.py 的区别（两份导出为什么不合并）：
  试卷是**结构化**的（题干/答案/图片几个固定字段），照着字段往 docx 里塞就行；
  笔记是**一篇 Markdown 正文**，得先解析出标题、段落、列表、表格、代码块、图片、公式，
  再逐块翻译成 docx 元素。两者的输入形态完全不同，硬合并只会互相拖累。

公式（本模块的重点）：
  正文里的 `$…$` / `$$…$$` 走 adapters/notes_math.py 转成 **OMML**，
  也就是 Word 的**原生公式** —— 导出后能双击编辑、转 PDF 排版正确。
  转换不可用时（没 node / 没装 Office）不报错，公式按 `$…$` 原文写上去，
  并在返回值里说明，绝不因为"公式渲染不了"让整份导出失败。

PDF 的生成方式：
  先出 docx，再交给 adapters/office.py 用 Word 导出 PDF。
  这样公式、字体、分页全部由 Word 负责，保真度最高；
  代价是 PDF 需要本机有 Word（试卷导出用的是 pymupdf 纯排版，没有这个依赖）。
  没有 Word 时抛 PdfUnavailable，路由层翻成 503 并给出可操作的提示。

板书（ink）：
  笔画是"覆盖在预览上的手写层"，它的坐标是相对**当时预览区域**归一化的 0~1，
  而预览高度随内容变化、这个比例没有存进库里。所以导出时把它渲染成一张独立的
  「板书」图（A4 比例画布），而不是强行叠在正文上 —— 叠上去必然会错位。
"""
from __future__ import annotations

import io
import re

from .. import config
from ..adapters import notes_math, office

# 正文里引用到的配图，形如 ![说明](/api/notes/files/nb_ab12.png)
_IMG_URL_RE = re.compile(r"^/api/notes/files/([A-Za-z0-9_.\-]+)$")

# A4 与页边距（和 paper_export 保持一致，打印出来观感统一）
A4_W, A4_H = 595.0, 842.0
MARGIN = 56.0
CONTENT_W = A4_W - MARGIN * 2

# 板书画布的假定宽度：笔画粗细是按屏幕像素存的（canvas 宽度约 794px），
# 折算到页面上要按比例缩一下，否则 10px 的笔在纸上会变成很粗的一条。
_INK_ASSUMED_CANVAS_W = 794.0


class PdfUnavailable(RuntimeError):
    """本机没有 Word，无法生成高保真 PDF。"""


# ================================================================ Markdown 解析
# 只实现笔记里真正会用到的子集：标题 / 段落 / 有序无序列表 / 引用 / 代码块 /
# 表格 / 分隔线 / 图片 / 公式。表格里塞完整 HTML 这种就不管了。
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_ORDERED_RE = re.compile(r"^(\s*)(\d{1,3})[.)]\s+(.*)$")
_BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_HR_RE = re.compile(r"^\s*([-*_])\s*(?:\1\s*){2,}$")
_IMG_ONLY_RE = re.compile(r"^\s*!\[([^\]]*)\]\(([^)\s]+)(?:\s+\"[^\"]*\")?\)\s*$")
_FENCE_RE = re.compile(r"^\s*```+\s*([\w+-]*)\s*$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$")


def _split_table_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


# ---- 行内解析：把一段文字切成若干 run（文字 / 加粗 / 斜体 / 代码 / 公式 / 链接）
_INLINE_RE = re.compile(
    r"(?P<code>`[^`]+`)"
    r"|(?P<bold>\*\*[^*]+\*\*|__[^_]+__)"
    r"|(?P<strike>~~[^~]+~~)"
    r"|(?P<italic>\*[^*\s][^*]*\*|_[^_\s][^_]*_)"
    r"|(?P<link>\[[^\]]*\]\([^)\s]+\))"
    r"|(?P<formula>\$(?!\$)[^$]+\$)"
)


def _parse_inline(text: str, formulas: list[str]) -> list[dict]:
    """切成 run 列表。公式同时登记进 formulas（后面统批转 OMML）。"""
    out: list[dict] = []
    pos = 0
    for m in _INLINE_RE.finditer(text or ""):
        if m.start() > pos:
            out.append({"t": "text", "v": text[pos:m.start()]})
        kind = m.lastgroup
        raw = m.group()
        if kind == "code":
            out.append({"t": "code", "v": raw[1:-1]})
        elif kind == "bold":
            out.append({"t": "bold", "v": raw[2:-2]})
        elif kind == "strike":
            out.append({"t": "strike", "v": raw[2:-2]})
        elif kind == "italic":
            out.append({"t": "italic", "v": raw[1:-1]})
        elif kind == "link":
            label, url = re.match(r"\[([^\]]*)\]\(([^)\s]+)\)", raw).groups()
            out.append({"t": "link", "v": label or url, "url": url})
        elif kind == "formula":
            formulas.append(raw[1:-1])
            out.append({"t": "formula", "i": len(formulas) - 1})
        pos = m.end()
    if pos < len(text or ""):
        out.append({"t": "text", "v": text[pos:]})
    return out


def _parse_blocks(md: str, formulas: list[str]) -> list[dict]:
    """Markdown -> 块列表。公式（行内与独立）登记进 formulas。"""
    lines = (md or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[dict] = []
    i = 0
    n = len(lines)

    def flush_para(buf: list[str]) -> None:
        if buf:
            blocks.append({"k": "para", "inl": _parse_inline(" ".join(buf).strip(), formulas)})

    para: list[str] = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 围栏代码块
        if _FENCE_RE.match(line):
            flush_para(para); para = []
            lang = _FENCE_RE.match(line).group(1)
            i += 1
            code: list[str] = []
            while i < n and not _FENCE_RE.match(lines[i]):
                code.append(lines[i])
                i += 1
            i += 1                       # 吃掉收尾的 ```
            blocks.append({"k": "code", "lang": lang, "text": "\n".join(code)})
            continue

        if not stripped:
            flush_para(para); para = []
            i += 1
            continue

        # 独立公式：$$ … $$（可能跨行）
        if stripped.startswith("$$"):
            flush_para(para); para = []
            buf = stripped[2:]
            if buf.endswith("$$") and len(buf) >= 0:
                tex = buf[:-2]
                i += 1
            else:
                parts = [buf]
                i += 1
                while i < n and "$$" not in lines[i]:
                    parts.append(lines[i])
                    i += 1
                if i < n:
                    parts.append(lines[i].split("$$")[0])
                    i += 1
                tex = "\n".join(parts)
            tex = tex.strip()
            if tex:
                formulas.append(tex)
                blocks.append({"k": "formula", "i": len(formulas) - 1, "tex": tex})
            continue

        m = _HEADING_RE.match(stripped)
        if m:
            flush_para(para); para = []
            blocks.append({"k": "heading", "level": len(m.group(1)),
                           "inl": _parse_inline(m.group(2).strip(), formulas)})
            i += 1
            continue

        if _HR_RE.match(stripped):
            flush_para(para); para = []
            blocks.append({"k": "hr"})
            i += 1
            continue

        # 整行一张图
        m = _IMG_ONLY_RE.match(line)
        if m:
            flush_para(para); para = []
            blocks.append({"k": "image", "alt": m.group(1), "url": m.group(2)})
            i += 1
            continue

        # 引用（连续的 > 行合成一段）
        if stripped.startswith(">"):
            flush_para(para); para = []
            quote: list[str] = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append({"k": "quote", "inl": _parse_inline(" ".join(quote).strip(), formulas)})
            continue

        # 表格：本行 + 下一行是分隔行才算
        if "|" in line and i + 1 < n and _TABLE_SEP_RE.match(lines[i + 1]):
            flush_para(para); para = []
            header = _split_table_row(line)
            i += 2
            rows: list[list[list[dict]]] = []
            while i < n and "|" in lines[i] and lines[i].strip():
                cells = _split_table_row(lines[i])
                rows.append([_parse_inline(c, formulas) for c in cells])
                i += 1
            blocks.append({"k": "table", "header": [_parse_inline(c, formulas) for c in header],
                           "rows": rows})
            continue

        # 列表（连续同类行合成一组；缩进深度按 2 空格算一层，最多 3 层）
        mo, mb = _ORDERED_RE.match(line), _BULLET_RE.match(line)
        if mo or mb:
            flush_para(para); para = []
            ordered = bool(mo)          # 以第一行为准 —— 混排的列表本来就不该出现
            items: list[dict] = []
            while i < n:
                if not lines[i].strip():
                    break
                cur_o, cur_b = _ORDERED_RE.match(lines[i]), _BULLET_RE.match(lines[i])
                if not (cur_o or cur_b):
                    break
                mm = cur_o or cur_b
                # ⚠️ 两类正则的捕获组不一样（有序是 缩进/序号/正文，无序是 缩进/正文），
                #    必须按"这一行到底匹配了哪个"取组号。曾经按组的 ordered 去取 group(3)，
                #    遇到 "- a" 后面跟 "1. b" 这种混排直接 IndexError 把整篇导出打断。
                indent = len(mm.group(1).replace("\t", "  "))
                text = mm.group(3) if cur_o else mm.group(2)
                items.append({"depth": min(3, indent // 2),
                              "inl": _parse_inline(text, formulas)})
                i += 1
            blocks.append({"k": "list", "ordered": ordered, "items": items})
            continue

        para.append(stripped)
        i += 1

    flush_para(para)
    return blocks


# ================================================================ 板书渲染
def _hex_to_rgb(s: str) -> tuple[float, float, float]:
    s = (s or "#000000").lstrip("#")
    if len(s) == 3:
        s = "".join(c * 2 for c in s)
    try:
        return (int(s[0:2], 16) / 255, int(s[2:4], 16) / 255, int(s[4:6], 16) / 255)
    except (ValueError, IndexError):
        return (0.0, 0.0, 0.0)


def render_ink_png(ink: list) -> bytes | None:
    """把笔画渲染成一张 PNG（A4 比例、透明底）。没有笔画返回 None。

    为什么用 A4 比例：笔画坐标是相对「当时预览区域」归一化的，而预览高度
    随内容长短变化、这个比例没存进库。用固定比例至少保证同一篇笔记每次导出
    的样子一致，也避免按 B 盒自适应导致圆形被拉成椭圆。
    """
    strokes = []
    for s in ink or []:
        if not isinstance(s, dict):
            continue
        pts = [p for p in (s.get("points") or []) if isinstance(p, (list, tuple)) and len(p) >= 2]
        if len(pts) < 2:
            continue
        strokes.append((_hex_to_rgb(s.get("color") or "#000000"),
                        float(s.get("width") or 3.0), pts))
    if not strokes:
        return None

    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=A4_W, height=A4_H)
    scale = A4_W / _INK_ASSUMED_CANVAS_W      # 屏幕 px -> 页面 pt
    for rgb, width, pts in strokes:
        page.draw_polyline(
            [pymupdf.Point(p[0] * A4_W, p[1] * A4_H) for p in pts],
            color=rgb,
            width=max(0.6, width * scale),
            lineCap=1,                        # 1 = 圆头，和画布上 ctx.lineCap='round' 一致
            lineJoin=1,
        )
    pm = page.get_pixmap(alpha=True, dpi=150)
    png = pm.tobytes("png")
    doc.close()
    return png


# ================================================================ DOCX
def build_docx(note: dict, include_ink: bool = True) -> tuple[bytes, dict]:
    """返回 (docx 字节, 情况说明)。说明里会写公式是渲染成了真公式还是降级成了原文。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    formulas: list[str] = []
    blocks = _parse_blocks(note.get("content") or "", formulas)

    # 所有公式一次批量转换（一次 node 调用），失败的项落到 None 走纯文本
    ommls = notes_math.latex_to_omml(formulas) if formulas else []
    converted = sum(1 for x in ommls if x)
    # 注入失败数：转换成功但塞进 docx 时出错（不该发生，但发生必须看得见 —— 
    # 曾经就是因为这里被 except 吞掉，整篇公式全成了 $…$ 原文而没人发现）
    inject_failed = 0
    info = {
        "formulas": len(formulas),
        "formulas_rendered": converted,
        "formula_note": "",
    }
    if formulas and (converted < len(formulas) or inject_failed):
        _, reason = notes_math.availability()
        parts = []
        if converted < len(formulas):
            parts.append(f"{len(formulas) - converted} 个公式未能转成 Word 公式（多数是 LaTeX 写法有误），已按 $…$ 原文写入")
        if inject_failed:
            parts.append(f"{inject_failed} 个公式写入 Word 时出错，已按原文写入")
        info["formula_note"] = "；".join(parts) + (
            f"。原因：{reason}" if reason and converted == 0 else "")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    content_cm = 21.0 - 2.0 * 2

    # python-docx 默认字体是西文的，中文会走回退。必须显式写 w:eastAsia。
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11.5)
    rf = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), "宋体")
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")

    # —— 标题与元信息
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run(note.get("title") or "未命名笔记")
    tr.bold = True
    tr.font.size = Pt(18)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(f"笔记导出 · {note.get('updated_at', '').replace('T', ' ')}")
    mr.font.size = Pt(9)
    mr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    def add_inline(paragraph, inl: list[dict], base_size: float = 11.5) -> None:
        """把 run 列表写进一个段落，公式转成真公式（OMML）内联塞进去。"""
        nonlocal inject_failed
        for r in inl or []:
            t = r.get("t")
            if t == "formula":
                omml = ommls[r["i"]] if r["i"] < len(ommls) else None
                if omml:
                    try:
                        paragraph._p.append(parse_xml(notes_math.omml_inline(omml)))
                        continue
                    except Exception:  # noqa: BLE001  注入失败退回原文，但必须记一笔
                        inject_failed += 1
                run = paragraph.add_run(f"${formulas[r['i']]}$")
                run.font.size = Pt(base_size)
                continue
            if t == "text":
                run = paragraph.add_run(r["v"])
                run.font.size = Pt(base_size)
                continue
            if t == "link":
                run = paragraph.add_run(r["v"])
                run.font.size = Pt(base_size)
                run.font.color.rgb = RGBColor(0x1E, 0x66, 0xC8)
                if r.get("url"):
                    tail = paragraph.add_run(f"（{r['url']}）")
                    tail.font.size = Pt(8.5)
                    tail.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
                continue
            run = paragraph.add_run(r["v"])
            run.font.size = Pt(base_size)
            if t == "bold":
                run.bold = True
            elif t == "italic":
                run.italic = True
            elif t == "strike":
                run.font.strike = True
            elif t == "code":
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)

    def add_picture(url: str, alt: str = "") -> bool:
        m = _IMG_URL_RE.match(url or "")
        if not m:
            return False
        path = config.NOTES_DIR / m.group(1)
        if not path.exists():
            return False
        # 按可用宽度放，同时避免特别长的竖图撑破一页
        try:
            import pymupdf
            with pymupdf.open(str(path)) as im:
                pw, ph = im[0].rect.width, im[0].rect.height
            ratio = (ph / pw) if pw else 1.0
        except Exception:  # noqa: BLE001
            ratio = 1.0
        width_cm = content_cm * 0.92
        max_h_cm = 29.7 - 2.0 * 2 - 2.0
        if width_cm * ratio > max_h_cm:
            width_cm = max_h_cm / max(ratio, 0.01)
        doc.add_picture(str(path), width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        return True

    def add_hr() -> None:
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(
            '<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:bottom w:val="single" w:sz="6" w:space="1" w:color="BBBBBB"/></w:pBdr>'
        ))

    # —— 逐块翻译
    for b in blocks:
        k = b["k"]
        if k == "heading":
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            add_inline(h, b["inl"], base_size=max(12.0, 19.0 - b["level"] * 1.6))
            for r in h.runs:
                r.bold = True
        elif k == "para":
            add_inline(doc.add_paragraph(), b["inl"])
        elif k == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            add_inline(p, b["inl"])
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif k == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            r = p.add_run(b["text"])
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        elif k == "list":
            for it in b["items"]:
                style = "List Number" if b["ordered"] else "List Bullet"
                if it["depth"]:
                    style += " %d" % min(3, it["depth"] + 1)
                try:
                    p = doc.add_paragraph(style=style)
                except KeyError:      # 模板里没有这个层级的样式就退回一级
                    p = doc.add_paragraph(style="List Number" if b["ordered"] else "List Bullet")
                add_inline(p, it["inl"])
        elif k == "table":
            cols = max(len(b["header"]), max((len(r) for r in b["rows"]), default=0))
            if cols:
                tbl = doc.add_table(rows=1, cols=cols)
                tbl.style = "Table Grid"
                for j, cell in enumerate(b["header"][:cols]):
                    para = tbl.rows[0].cells[j].paragraphs[0]
                    add_inline(para, cell, base_size=10.5)
                    for r in para.runs:
                        r.bold = True
                for row in b["rows"]:
                    cells = tbl.add_row().cells
                    for j, cell in enumerate(row[:cols]):
                        add_inline(cells[j].paragraphs[0], cell, base_size=10.5)
                doc.add_paragraph()
        elif k == "image":
            if not add_picture(b["url"], b.get("alt", "")):
                p = doc.add_paragraph()
                r = p.add_run(f"[图片：{b.get('alt') or b['url']}]")
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        elif k == "formula":
            omml = ommls[b["i"]] if b["i"] < len(ommls) else None
            ok = False
            if omml:
                try:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p._p.append(parse_xml(notes_math.omml_display(omml)))
                    ok = True
                except Exception:  # noqa: BLE001
                    inject_failed += 1
                    ok = False
            if not ok:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(f"$${b['tex']}$$")
                r.font.size = Pt(11.5)
        elif k == "hr":
            add_hr()

    # —— 导出情况说明（必须放在正文建完之后：注入失败数是构建过程中累加出来的）
    if formulas and (converted < len(formulas) or inject_failed):
        _, reason = notes_math.availability()
        parts = []
        if converted < len(formulas):
            parts.append(
                f"{len(formulas) - converted} 个公式未能转成 Word 公式"
                "（多数是 LaTeX 写法有误），已按 $…$ 原文写入"
            )
        if inject_failed:
            parts.append(f"{inject_failed} 个公式写入 Word 时出错，已按原文写入")
        info["formulas_injected"] = converted - inject_failed
        info["formula_note"] = "；".join(parts) + (
            f"。原因：{reason}" if reason and converted == 0 else "")

    # —— 板书（手写标注）附在最后
    if include_ink:
        png = render_ink_png(note.get("ink") or [])
        if png:
            doc.add_page_break()
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = cap.add_run("板书（原笔记上的手写标注，按原始相对位置绘制）")
            cr.font.size = Pt(10)
            cr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            doc.add_picture(io.BytesIO(png), width=Cm(content_cm * 0.86))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), info


# ================================================================ PDF
def build_pdf(note: dict, include_ink: bool = True) -> tuple[bytes, dict]:
    """docx -> Word -> PDF。没有 Word 时抛 PdfUnavailable。"""
    docx_bytes, info = build_docx(note, include_ink=include_ink)
    try:
        return office.docx_bytes_to_pdf(docx_bytes), info
    except office.OfficeUnavailable as e:
        raise PdfUnavailable(
            "导出 PDF 需要本机装有 Microsoft Word（服务端用 Word 排版，公式才能正确显示）。"
            f"当前不可用：{e}。可以先「导出 Word」，再用 Word 另存为 PDF。"
        ) from e
    except office.WordConvertError as e:
        raise PdfUnavailable(str(e)) from e


def safe_filename(title: str, ext: str) -> str:
    """下载文件名：去掉路径分隔符等非法字符。"""
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", (title or "笔记")).strip(" .") or "笔记"
    return f"{name[:60]}.{ext}"
