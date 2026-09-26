# -*- coding: utf-8 -*-
"""课后反馈导出：txt / Word(.docx) / PDF。

三种格式的分工（以及为什么这么分）：
  · **txt** —— 最通用、任何设备都能看（家长微信里发一段文字最省事）。
    按需求，**txt 不带图片**：纯文本里塞图片路径只会变成一堆看不懂的字符。
  · **docx** —— 带配图，家长可打印、老师可再改。
  · **pdf** —— 先出 docx，再交给 adapters/office.py 用本机 Word 导出（版式最准）。
    没有 Word 时抛 PdfUnavailable，由路由翻成 503 并给可照做的提示。

内容组织：整篇正文（AI 按模板整理出的成品）**优先**，没有就退回四段式
（课堂表现 / 存在问题 / 作业布置 / 下次安排）+ 综合状态 + 能力评分 + 配图。
空字段**不输出标题**——留一堆「（空）」比不写更难看，这是给家长看的正式文本。
"""
from __future__ import annotations

import io
import re

from .. import config
from ..adapters import office

# 反馈导出的四个字段：(键, 中文标题)
SECTIONS = [
    ("performance", "课堂表现"),
    ("problems", "存在问题"),
    ("homework", "作业布置"),
    ("next_plan", "下次安排"),
]

# 图片 URL 形如 /api/feedbacks/files/fb_ab12.png
_IMG_URL_RE = re.compile(r"^/api/feedbacks/files/([A-Za-z0-9_.\-]+)$")

# 整篇正文里的栏目标题行：「【本次课堂内容】」
_DOC_HEAD_RE = re.compile(r"^【.+】$")

STAR_FULL, STAR_EMPTY = "★", "☆"

MAX_SCORE = 5          # 能力维度满分（与前端评分控件的 1-5 对应）

# 雷达图配色：和前端 components/AbilityRadar.js 保持一致 ——
# 同一个学生、同一份数据，屏幕上和导出文件里应该是同一个样子。
_R_CUR = (0.145, 0.388, 0.921)     # #2563eb 本次
_R_PREV = (0.6, 0.63, 0.7)         # #98a2b3 上次
_R_GRID = (0.894, 0.906, 0.925)    # #e4e7ec 网格
_R_TEXT = (0.357, 0.392, 0.447)    # #5b6472 文字


class PdfUnavailable(RuntimeError):
    """本机没有 Word，无法生成高保真 PDF。"""


def _stars(rating: int | None) -> str:
    if not rating:
        return ""
    n = max(0, min(5, int(rating)))
    return STAR_FULL * n + STAR_EMPTY * (5 - n)


def _meta_line(ctx: dict) -> str:
    """标题下面那行：学生 / 时间 / 课程 / 状态。缺失的项直接不出现。"""
    parts = []
    if ctx.get("student_name"):
        parts.append(f"学生：{ctx['student_name']}")
    if ctx.get("when"):
        parts.append(f"时间：{ctx['when']}")
    if ctx.get("topic"):
        parts.append(f"内容：{ctx['topic']}")
    st = _stars(ctx.get("rating"))
    if st:
        parts.append(f"综合状态：{st}")
    return "　　".join(parts)


def _score_lines(scores: list[dict], prev: dict[int, int] | None = None) -> list[str]:
    """「· 概念理解：4 / 5（上次 3）」。

    带上「上次」很有必要：txt 版本没有雷达图，光列本次分数看不出变化。
    """
    out = []
    for s in scores or []:
        if not s.get("score"):
            continue
        name = s.get("name") or ("#" + str(s.get("dim_id")))
        line = f"{name}：{s['score']} / {MAX_SCORE}"
        before = (prev or {}).get(s.get("dim_id"))
        if before:
            line += f"（上次 {before}）"
        out.append(line)
    return out


def build_radar_png(dims: list[dict], size: int = 660, dpi: int = 200) -> bytes | None:
    """手绘能力雷达图 → PNG 字节；画不了就返回 None。

    dims: [{"name": "概念理解", "latest": 4, "previous": 3}]，previous 可为 None。

    为什么手绘而不是引图表库：本项目「零新依赖」是硬约束（venv 里没有 matplotlib），
    而 PyMuPDF 本来就在用（探测图片尺寸），它能画线、能写中文（内置 china-s 字体）、
    能导出 PNG —— 三样刚好都满足。

    为什么必须有「上次」那条：一张孤立的雷达图家长看不出好坏，两条叠在一起
    「变化」才看得见 —— 和界面上的 AbilityRadar 是同一个道理。
    """
    usable = [d for d in (dims or []) if d.get("latest")]
    if len(usable) < 3:
        return None          # 两个点的「雷达图」没有意义，不如只列数字
    try:
        import math

        import pymupdf as fitz

        font = fitz.Font("china-s")      # MuPDF 内置简体中文字体，不依赖系统装了什么
        fs = 13
        labels = [f"{d['name']} {d['latest']:g}" for d in usable]
        # 半径按最宽的标签算，否则长维度名会顶出画布被截断（实测过）
        max_w = max(font.text_length(t, fontsize=fs) for t in labels)
        R = size / 2 - max_w - 40
        if R < 40:
            return None

        page = fitz.open().new_page(width=size, height=size)
        cx = cy = size / 2
        n = len(usable)
        has_prev = any(d.get("previous") for d in usable)

        def pt(i: int, v: float):
            a = -math.pi / 2 + 2 * math.pi * i / n
            r = R * max(0.0, min(float(v), MAX_SCORE)) / MAX_SCORE
            return fitz.Point(cx + r * math.cos(a), cy + r * math.sin(a))

        # 网格：每 1 分一圈 + 从圆心发散的轴线
        sh = page.new_shape()
        for lvl in range(1, MAX_SCORE + 1):
            sh.draw_polyline([pt(i, lvl) for i in range(n)])
            sh.finish(color=_R_GRID, width=0.8, closePath=True)
        for i in range(n):
            sh.draw_line(fitz.Point(cx, cy), pt(i, MAX_SCORE))
        sh.finish(color=_R_GRID, width=0.8)
        sh.commit()

        # 上次（灰虚线）先画，压在下面
        if has_prev:
            sh = page.new_shape()
            sh.draw_polyline([pt(i, d.get("previous") or 0) for i, d in enumerate(usable)])
            sh.finish(color=_R_PREV, fill=_R_PREV, fill_opacity=0.16,
                      width=1.2, dashes="4 3", closePath=True)
            sh.commit()

        # 本次（蓝实线）
        sh = page.new_shape()
        sh.draw_polyline([pt(i, d["latest"]) for i, d in enumerate(usable)])
        sh.finish(color=_R_CUR, fill=_R_CUR, fill_opacity=0.16, width=1.6, closePath=True)
        sh.commit()

        # 标签：PyMuPDF 没有 text-anchor，得自己按角度算对齐
        for i, text in enumerate(labels):
            a = -math.pi / 2 + 2 * math.pi * i / n
            cos = math.cos(a)
            lx = cx + (R + 30) * cos
            ly = cy + (R + 30) * math.sin(a) + fs * 0.35
            w = font.text_length(text, fontsize=fs)
            if cos > 0.3:
                x = lx                      # 右半边：左对齐
            elif cos < -0.3:
                x = lx - w                  # 左半边：右对齐
            else:
                x = lx - w / 2              # 顶/底：居中
            page.insert_text((x, ly), text, fontsize=fs, fontname="china-s", color=_R_TEXT)

        ly = size - 22
        page.insert_text((cx - 78, ly), "■ 本次", fontsize=12, fontname="china-s", color=_R_CUR)
        if has_prev:
            page.insert_text((cx + 4, ly), "▨ 上次", fontsize=12, fontname="china-s", color=_R_PREV)

        png = page.get_pixmap(dpi=dpi, alpha=False).tobytes("png")
        page.parent.close()
        return png
    except Exception:  # noqa: BLE001
        # 画不出来就退回文字列表 —— 绝不能因为一张图让整份反馈导不出来
        return None


def _image_paths(images: list[str]) -> list:
    """把 URL 换成磁盘文件。只接受本项目自己的 /api/feedbacks/files/ 形式，防路径穿越。"""
    out = []
    for url in images or []:
        m = _IMG_URL_RE.match(url or "")
        if not m:
            continue
        p = config.FEEDBACK_DIR / m.group(1)
        if p.exists() and p.is_file():
            out.append(p)
    return out


# ---------------------------------------------------------------- TXT
def build_txt(fb: dict, ctx: dict, doc: str | None = None) -> bytes:
    """纯文本。刻意不带图片（需求明确：txt 不用图片）。

    doc 是整篇正文（AI 整理的成品）。有它就整篇导出 —— 那是老师确认过的正式文本；
    没有才退回四段式。不把两者混在一起拼：同一件事写两遍，家长会觉得乱。
    """
    if (doc or "").strip():
        lines: list[str] = [doc.strip(), ""]
        scores_doc = _score_lines(fb.get("ability_scores") or [], ctx.get("ability_prev"))
        if scores_doc:
            lines.append("【能力评分】")
            lines.extend("· " + s for s in scores_doc)
            lines.append("")
        images_doc = fb.get("images") or []
        if images_doc:
            lines.append(f"（本篇附有 {len(images_doc)} 张图片，请查看 Word 或 PDF 版本）")
            lines.append("")
        text_doc = "\n".join(lines).rstrip() + "\n"
        return text_doc.encode("utf-8-sig")

    lines: list[str] = ["课后反馈", "=" * 24]
    meta = _meta_line(ctx)
    if meta:
        lines.append(meta)
    lines.append("")

    for key, title in SECTIONS:
        text = (fb.get(key) or "").strip()
        if not text:
            continue
        lines.append(f"【{title}】")
        lines.extend(text.splitlines())
        lines.append("")

    scores = _score_lines(fb.get("ability_scores") or [], ctx.get("ability_prev"))
    if scores:
        lines.append("【能力评分】")
        lines.extend("· " + s for s in scores)
        lines.append("")

    images = fb.get("images") or []
    if images:
        # 说清楚「有几张图、要去 Word/PDF 里看」，而不是默默丢掉
        lines.append(f"（本篇附有 {len(images)} 张图片，请查看 Word 或 PDF 版本）")
        lines.append("")

    text = "\n".join(lines).rstrip() + "\n"
    # utf-8-sig：带 BOM。不加的话 Windows 记事本、Excel 打开中文可能乱码 ——
    # 这个文件是要发给家长/贴在报告里的，不能出现「打开是乱码」这种事。
    return text.encode("utf-8-sig")


# ---------------------------------------------------------------- DOCX / PDF
def _add_body(document, text: str, Pt) -> None:
    """把一段正文写进 docx：整行是【栏目】的当作小标题加粗，其余按段落写。

    整篇正文是 AI 按老师给的模板写的，栏目名不固定（可能是【易错内容】
    【作业预计时长】这类四段里没有的），所以这里按「长什么样」识别，
    而不是去比对一张写死的栏目表。
    """
    for line in (text or "").splitlines():
        s = line.strip()
        if not s:
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if _DOC_HEAD_RE.match(s):
            p.paragraph_format.space_before = Pt(10)
            r = p.add_run(s)
            r.bold = True
            r.font.size = Pt(12)
        else:
            p.add_run(s)


def build_docx(fb: dict, ctx: dict, doc: str | None = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = (doc or "").strip()
    whole = bool(doc)

    document = Document()
    sec = document.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    content_cm = 21.0 - 2.0 * 2

    # python-docx 默认西文字体，中文会走回退。必须显式写 w:eastAsia。
    normal = document.styles["Normal"]
    normal.font.size = Pt(11.5)
    rf = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), "宋体")
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")

    if whole:
        # 整篇正文通常自带抬头（学生名-日期 课堂反馈 / 科目 / 上课时间…），
        # 所以不再另加标题，否则会变成「课后反馈」+ 学生自己那行抬头两个标题。
        _add_body(document, doc, Pt)
        _append_docx_extras(document, fb, ctx, Pt, WD_ALIGN_PARAGRAPH, Cm)
        buf_whole = io.BytesIO()
        document.save(buf_whole)
        return buf_whole.getvalue()

    tp = document.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run("课后反馈")
    tr.bold = True
    tr.font.size = Pt(18)

    meta = _meta_line(ctx)
    if meta:
        mp = document.add_paragraph()
        mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mr = mp.add_run(meta)
        mr.font.size = Pt(10)
        mr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    for key, title in SECTIONS:
        text = (fb.get(key) or "").strip()
        if not text:
            continue
        h = document.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
        hr = h.add_run(f"【{title}】")
        hr.bold = True
        hr.font.size = Pt(12)
        for line in text.splitlines():
            p = document.add_paragraph(line)
            p.paragraph_format.space_after = Pt(2)

    _append_docx_extras(document, fb, ctx, Pt, WD_ALIGN_PARAGRAPH, Cm)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _append_docx_extras(document, fb: dict, ctx: dict, Pt, WD_ALIGN_PARAGRAPH, Cm) -> None:
    """两种模式共用的尾巴：能力评分（雷达图）+ 附图。"""
    scores = _score_lines(fb.get("ability_scores") or [], ctx.get("ability_prev"))
    radar = build_radar_png(ctx.get("ability_radar") or [])
    if scores or radar:
        h = document.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
        hr = h.add_run("【能力评分】")
        hr.bold = True
        hr.font.size = Pt(12)
        if radar:
            # 图里的标签已经带了数值，不再重复罗列文字 —— 一份给家长的报告
            # 同一件事说两遍反而显得啰嗦
            document.add_picture(io.BytesIO(radar), width=Cm(11.5))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            for s in scores:
                p = document.add_paragraph("· " + s)
                p.paragraph_format.space_after = Pt(0)

    pics = _image_paths(fb.get("images") or [])
    if not pics:
        return
    h = document.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(2)
    hr = h.add_run("【附图】")
    hr.bold = True
    hr.font.size = Pt(12)
    content_cm = 21.0 - 2.0 * 2
    for path in pics:
        # 按可用宽度放，同时避免竖长图撑破一页
        try:
            import pymupdf
            with pymupdf.open(str(path)) as im:
                pw, ph = im[0].rect.width, im[0].rect.height
            ratio = (ph / pw) if pw else 1.0
        except Exception:  # noqa: BLE001  图片坏了不该让整篇导不出来
            ratio = 1.0
        width_cm = content_cm * 0.9
        max_h = 29.7 - 2.0 * 2 - 3.0
        if width_cm * ratio > max_h:
            width_cm = max_h / max(ratio, 0.01)
        document.add_picture(str(path), width=Cm(width_cm))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_pdf(fb: dict, ctx: dict, doc: str | None = None) -> bytes:
    """docx -> Word -> PDF。没有 Word 时抛 PdfUnavailable。"""
    docx_bytes = build_docx(fb, ctx, doc)
    try:
        return office.docx_bytes_to_pdf(docx_bytes)
    except office.OfficeUnavailable as e:
        raise PdfUnavailable(
            "导出 PDF 需要本机装有 Microsoft Word（服务端用 Word 排版）。"
            f"当前不可用：{e}。可以先「导出 Word」，再用 Word 另存为 PDF。"
        ) from e
    except office.WordConvertError as e:
        raise PdfUnavailable(str(e)) from e


def safe_filename(ctx: dict, ext: str) -> str:
    """文件名带上学生和日期，家长收到一眼能认出来。"""
    name = ctx.get("student_name") or "学生"
    when = (ctx.get("when") or "")[:10]
    base = f"课后反馈-{name}-{when}" if when else f"课后反馈-{name}"
    base = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", base).strip(" .") or "课后反馈"
    return f"{base[:60]}.{ext}"
