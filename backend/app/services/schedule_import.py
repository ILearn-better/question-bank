# -*- coding: utf-8 -*-
"""课表 Excel 导入 —— 把老师手里的排课表（.xlsx）解析成结构化的课时行。

设计要点（改这块之前先读）
------------------------------------------------
1. **解析层永不抛异常**。Excel 是手写出来的，脏数据是常态：日期写成「10/1」、
   时间写成全角冒号的「19：00」、时长写成「1小时」。这些都要变成「这一行有什么问题」
   的**清单**交给老师看，而不是一个 500 或者一句「导入失败」。
2. **只认表头，不认列序号**。老师可能调换列顺序、插一列自己的「备注」。
   所以先按第一行（或前几行里最像表头的那行）建立「字段 → 列」的映射，再逐行取值。
   缺了必填列就整份拒绝 —— 继续解析没有意义，老师需要先改表。
3. **不做「智能猜测」**。学生名匹配不上、状态写法不认识 —— 如实报告，绝不猜。
   猜错了会静默产生错数据，比直接报错糟得多（这个项目一贯的取舍，见 README
   「已知限制」里那条「宁可不做，不做错」）。
4. **本文件不碰数据库**。学生匹配与查重放在 router 里做 —— 这样解析逻辑可以
   脱离服务单测（test_schedule_import.py 就是这么测的）。

列的定义只有一处（下面的 COLUMNS）：模板生成、解析、文档说明都从它取，
避免「文档说能写 A，代码只认 B」这类不一致。
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, time, timedelta

try:                                    # openpyxl 缺失时不要让整个服务起不来
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except ImportError:                     # pragma: no cover
    openpyxl = None
    _HAS_OPENPYXL = False

MAX_ROWS = 2000                         # 一份课表最多这么多行，挡住误传大文件
HEADER_SCAN = 10                        # 表头可能不在第 1 行（上面有标题），往下扫几行


# ---------------------------------------------------------------- 列定义
# (字段名, 显示名, 是否必填, 别名, 模板/文档里的填写说明)
COLUMNS: list[tuple[str, str, bool, tuple[str, ...], str]] = [
    ("date", "日期", True, ("日期", "上课日期", "date", "上课日"),
     "2026-10-01。也可以直接写成 Excel 的日期格式"),
    ("time", "开始时间", True, ("开始时间", "时间", "上课时间", "time", "开始"),
     "19:00。写成「19:00-20:00」会顺便把时长算出来"),
    ("student", "学生", True, ("学生", "学生姓名", "姓名", "student", "名字"),
     "系统里已有的学生姓名，一个字都不能差"),
    ("duration_min", "时长（分钟）", False, ("时长（分钟）", "时长(分钟)", "时长", "分钟", "分钟数", "duration", "duration_min"),
     "留空按 60 分钟算"),
    ("topic", "本次内容", False, ("本次内容", "内容", "课程内容", "topic", "主题"),
     "如：三角函数图像变换"),
    ("mode", "上课形式", False, ("上课形式", "形式", "mode", "方式"),
     "线下 / 线上。留空按线下算"),
    ("location", "地点", False, ("地点", "location", "地址"),
     "线上课可以留空"),
    ("rate", "单价", False, ("单价", "课时费", "rate", "价格"),
     "这一节课的单价。留空就用学生档案里的默认单价"),
    ("status", "状态", False, ("状态", "status"),
     "已排课 / 已完成 / 补课 / 请假 / 已取消 / 已调课。留空按「已排课」算"),
]

FIELD_LABEL = {f: label for f, label, _req, _al, _hint in COLUMNS}
REQUIRED_FIELDS = [f for f, _l, req, _a, _h in COLUMNS if req]

# 状态 / 形式的写法对照。**只做写法归一，不做业务判断** ——
# 写「待上」能懂，写「看情况」就不懂，宁可报错让老师改。
STATUS_ALIASES = {
    "已排课": "scheduled", "排课": "scheduled", "待上": "scheduled", "计划": "scheduled",
    "scheduled": "scheduled",
    "已完成": "done", "完成": "done", "已上": "done", "上完": "done", "done": "done",
    "补课": "makeup", "makeup": "makeup",
    "请假": "leave", "leave": "leave",
    "已取消": "cancelled", "取消": "cancelled", "cancelled": "cancelled", "canceled": "cancelled",
    "已调课": "moved", "调课": "moved", "moved": "moved",
}
MODE_ALIASES = {
    "线下": "offline", "面授": "offline", "到店": "offline", "offline": "offline",
    "线上": "online", "网课": "online", "线上课": "online", "online": "online",
}


# ---------------------------------------------------------------- 小工具
def _norm(s) -> str:
    """表头/枚举值的归一化：去空白、全角括号与冒号转半角、去掉 `*` 和不换行空格。"""
    if s is None:
        return ""
    t = str(s).strip()
    t = t.replace("\u3000", "").replace("\xa0", "")
    t = t.replace("（", "(").replace("）", ")")
    t = t.replace("：", ":")
    t = re.sub(r"\s+", "", t)
    return t.replace("*", "").lower()


def _is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _s(v) -> str:
    """把单元格值变成干净的字符串（None → ''，数字去掉多余的 .0）。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_date(v, default_year: int | None = None) -> tuple[str | None, bool]:
    """→ (YYYY-MM-DD | None, 是否替用户补了年份)

    认得：Excel 日期单元格、'2026-10-01'、'2026/10/1'、'2026年10月1日'、
    '10-01'（缺年份 → 补 default_year 并标记）、Excel 序列号。
    """
    if _is_blank(v):
        return None, False
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d"), False
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d"), False
    if isinstance(v, (int, float)):
        # Excel 序列号（1900 纪元）。1900-02-29 那个历史 bug 用 1899-12-30 作原点即可对齐。
        try:
            return (date(1899, 12, 30) + timedelta(days=int(v))).strftime("%Y-%m-%d"), False
        except (OverflowError, ValueError):
            return None, False

    s = str(v).strip()
    s = s.replace("年", "-").replace("月", "-").replace("日", "").replace("/", "-").replace(".", "-")
    s = re.sub(r"\s+", "", s)
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", s)
    if m:
        y, mo, d = (int(x) for x in m.groups())
        try:
            return date(y, mo, d).strftime("%Y-%m-%d"), False
        except ValueError:
            return None, False
    m = re.match(r"^(\d{1,2})-(\d{1,2})$", s)
    if m and default_year:
        mo, d = (int(x) for x in m.groups())
        try:
            return date(default_year, mo, d).strftime("%Y-%m-%d"), True
        except ValueError:
            return None, False
    # 只剩月日数字连写，如 '1001' → 10-01
    m = re.match(r"^(\d{4})$", s)
    if m and default_year:
        mo, d = int(s[:2]), int(s[2:])
        try:
            return date(default_year, mo, d).strftime("%Y-%m-%d"), True
        except ValueError:
            return None, False
    return None, False


def parse_time(v) -> tuple[str | None, int | None]:
    """→ ('HH:MM' | None, 从时间区间推出的时长分钟 | None)

    认得：Excel 时间单元格、'19:00'、'19：00'（全角）、'1900'、
    '19:00-20:00' / '19:00~20:30'（**顺便把时长算出来**，省得老师写两列）。
    """
    if _is_blank(v):
        return None, None
    if isinstance(v, datetime):
        return v.strftime("%H:%M"), None
    if isinstance(v, time):
        return v.strftime("%H:%M"), None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        # 0.79 ≈ 19:00（Excel 的时间就是一天的小数部分）
        if 0 <= float(v) < 1:
            total = round(float(v) * 24 * 60)
            return f"{total // 60:02d}:{total % 60:02d}", None
        s = str(int(v))
    else:
        s = str(v).strip()

    s = s.replace("：", ":").replace("－", "-").replace("—", "-").replace("～", "~")
    s = re.sub(r"\s+", "", s)
    s = re.sub(r"^(上午|下午|晚上|中午)", "", s)
    # 老师常把星期也写进时间格（「周三 19:00」），剥掉它不影响时间本身
    s = re.sub(r"^(周|星期|礼拜)[一二三四五六日天][、,]?", "", s)

    # 时间区间：19:00-20:30 → 取开始 + 算时长
    m = re.match(r"^(\d{1,2}):(\d{1,2})\s*[-~至到]\s*(\d{1,2}):(\d{1,2})$", s)
    if m:
        h1, m1, h2, m2 = (int(x) for x in m.groups())
        if h1 > 23 or m1 > 59 or h2 > 24 or m2 > 59:
            return None, None
        mins = (h2 * 60 + m2) - (h1 * 60 + m1)
        if mins <= 0:
            mins += 24 * 60               # 跨零点（21:00-00:30）
        return f"{h1:02d}:{m1:02d}", mins

    # 只有开始时间
    m = re.match(r"^(\d{1,2}):(\d{1,2})$", s)
    if m:
        h, mi = int(m.group(1)), int(m.group(2))
        if h <= 23 and mi <= 59:
            return f"{h:02d}:{mi:02d}", None
        return None, None

    # '1900' / '930' 这种连写
    m = re.match(r"^(\d{3,4})$", s)
    if m:
        t = s.zfill(4)
        h, mi = int(t[:2]), int(t[2:])
        if h <= 23 and mi <= 59:
            return f"{h:02d}:{mi:02d}", None
    return None, None


# 「一」到「十」加「半」—— 只覆盖老师口头会写的这几种，不做通用的中文数字解析
# （「一又三分之二小时」这种，宁可不认，也不猜）。
_CN_NUM = {"半": 0.5, "一": 1, "两": 2, "二": 2, "三": 3, "四": 4, "五": 5,
           "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}


def parse_duration(v) -> tuple[int | None, str | None]:
    """→ (分钟 | None, 报错文案 | None)。认得 '90'、'90分钟'、'1.5小时'、'1小时30分'、'一小时'。"""
    if _is_blank(v):
        return None, None
    if isinstance(v, bool):
        return None, "时长填的不是数字"
    if isinstance(v, (int, float)):
        n = int(v)
        return (n, None) if 0 < n <= 24 * 60 else (None, f"时长 {v} 不像分钟数")

    s = str(v).strip()
    s = s.replace("：", ":").replace("　", "")
    if re.match(r"^\d+(\.\d+)?$", s):
        n = int(float(s))
        return (n, None) if 0 < n <= 24 * 60 else (None, f"时长 {v} 不像分钟数")

    # 1小时30分 / 1.5小时 / 90分钟 / 一小时 / 半小时
    num = r"(\d+(?:\.\d+)?|[一二两三四五六七八九十半])"
    h = re.search(num + r"\s*(?:个)?(?:小时|h|hour)", s, re.I)
    mi = re.search(num + r"\s*(?:分钟|分|min|m)\b", s, re.I)

    def _val(m):
        g = m.group(1)
        return _CN_NUM[g] if g in _CN_NUM else float(g)

    if h or mi:
        total = 0.0
        if h:
            total += _val(h) * 60
        if mi:
            total += _val(mi)
        total = int(round(total))
        return (total, None) if 0 < total <= 24 * 60 else (None, f"时长 {v} 不像分钟数")
    return None, f"看不懂的时长写法「{v}」（写分钟数，或「1.5小时」）"


def parse_money(v) -> tuple[float | None, str | None]:
    """→ (金额 | None, 报错文案 | None)。容忍 '¥300'、'300 元'、'300/小时'。"""
    if _is_blank(v):
        return None, None
    if isinstance(v, bool):
        return None, "单价填的不是数字"
    if isinstance(v, (int, float)):
        return (float(v), None) if v >= 0 else (None, f"单价 {v} 是负数")
    s = str(v).strip().replace("¥", "").replace("￥", "").replace("元", "")
    s = s.split("/")[0].split("每")[0].strip()
    try:
        n = float(s)
        return (n, None) if n >= 0 else (None, f"单价「{v}」是负数")
    except ValueError:
        return None, f"看不懂的单价「{v}」"


def parse_enum(v, aliases: dict[str, str], label: str) -> tuple[str | None, str | None]:
    if _is_blank(v):
        return None, None
    key = _norm(v)
    if key in aliases:
        return aliases[key], None
    opts = "、".join(sorted({k for k in aliases if not k.isascii()}))
    return None, f"看不懂的{label}「{v}」（可写：{opts}）"


# ---------------------------------------------------------------- 表头定位
def _match_columns(values: list) -> dict[str, int]:
    """把一行单元格值映射成 {字段名: 列下标(0-based)}。按别名归一化后**全等**匹配。"""
    out: dict[str, int] = {}
    for idx, raw in enumerate(values):
        key = _norm(raw)
        if not key:
            continue
        for field, _label, _req, aliases, _hint in COLUMNS:
            if field in out:
                continue
            if key in {_norm(a) for a in aliases}:
                out[field] = idx
                break
    return out


def _locate_header(ws) -> tuple[int, list, dict[str, int]]:
    """找表头行。老师常在第 1 行写个标题（「10 月课表」），所以往下扫几行，
    取「认出的列最多」的那一行；一行都没认出来就退回第 1 行。"""
    best = (1, [], {})
    best_hit = 0
    limit = min(ws.max_row or 1, HEADER_SCAN)
    for r in range(1, limit + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, (ws.max_column or 1) + 1)]
        m = _match_columns(vals)
        if len(m) > best_hit:
            best, best_hit = (r, vals, m), len(m)
        if best_hit >= len(COLUMNS):
            break
    return best


# ---------------------------------------------------------------- 主解析
def _empty_result(**kw) -> dict:
    base = {
        "ok": False,
        "error": None,
        "sheet": None,
        "sheets": [],
        "missing_columns": [],
        "matched_columns": {},
        "ignored_columns": [],
        "rows": [],
        "counts": {"total": 0, "ok": 0, "problem": 0},
        "notes": [],
    }
    base.update(kw)
    return base


def parse_workbook(data: bytes, default_year: int | None = None) -> dict:
    """解析一份 .xlsx → 结果 dict（**永不抛异常**，出了问题就放进 error / rows[].errors）。"""
    if not _HAS_OPENPYXL:
        return _empty_result(error="服务端没装 openpyxl，读不了 Excel。请先安装依赖后重启服务。")
    if not data:
        return _empty_result(error="文件是空的")
    default_year = default_year or datetime.now().year

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:                          # 不是 xlsx / 文件损坏 / 是老 .xls
        return _empty_result(
            error=f"读不了这个文件（{type(e).__name__}）。请确认是 .xlsx 格式 —— "
                  f"老的 .xls 需要在 Excel 里「另存为」成 .xlsx 再导入。"
        )

    try:
        sheets = list(wb.sheetnames)
        if not sheets:
            return _empty_result(error="这个文件里没有工作表")

        # 挑工作表：优先「必填列认全了」的第一个；都没有就退回第一个（好把缺什么告诉老师）
        chosen = None
        for name in sheets:
            ws = wb[name]
            hrow, hvals, mapping = _locate_header(ws)
            missing = [f for f in REQUIRED_FIELDS if f not in mapping]
            if not missing:
                chosen = (name, ws, hrow, hvals, mapping, missing)
                break
            if chosen is None:
                chosen = (name, ws, hrow, hvals, mapping, missing)
        name, ws, hrow, hvals, mapping, missing = chosen

        result = _empty_result(sheet=name, sheets=sheets)
        result["matched_columns"] = {
            f: _s(hvals[mapping[f]]) for f in mapping if mapping[f] < len(hvals)
        }
        result["ignored_columns"] = [
            _s(v) for i, v in enumerate(hvals)
            if _s(v) and i not in set(mapping.values())
        ]

        if missing:
            result["missing_columns"] = [FIELD_LABEL[f] for f in missing]
            result["error"] = (
                "表头里找不到这几列：" + "、".join(FIELD_LABEL[f] for f in missing) +
                "。请照「下载模板」里的表头填（列的顺序可以不一样，名字要对）。"
            )
            return result

        # ---- 逐行读 ----
        rows: list[dict] = []
        notes: list[str] = []
        filled_year = 0
        from_range = 0
        for r in range(hrow + 1, (ws.max_row or hrow) + 1):
            vals = [ws.cell(row=r, column=c).value for c in range(1, (ws.max_column or 1) + 1)]
            if all(_is_blank(v) for v in vals):
                continue                        # 整行空 —— 跳过（老师爱留空行分组）
            def cell(field):
                i = mapping.get(field)
                return vals[i] if i is not None and i < len(vals) else None

            errors: list[str] = []
            warnings: list[str] = []

            d, year_added = parse_date(cell("date"), default_year)
            if year_added:
                filled_year += 1
            if d is None:
                errors.append("日期看不懂" + (f"（「{_s(cell('date'))}」）" if _s(cell("date")) else "（空）"))

            t, mins_from_range = parse_time(cell("time"))
            if t is None:
                errors.append("开始时间看不懂" + (f"（「{_s(cell('time'))}」）" if _s(cell("time")) else "（空）"))

            student = _s(cell("student"))
            if not student:
                errors.append("没写学生")

            dur, derr = parse_duration(cell("duration_min"))
            if derr:
                errors.append(derr)
            if dur is None and mins_from_range:
                dur = mins_from_range            # 「19:00-20:00」推出来的
                from_range += 1
            if dur is None:
                dur = 60                         # 与排课弹窗的默认值保持一致

            mode, merr = parse_enum(cell("mode"), MODE_ALIASES, "上课形式")
            if merr:
                errors.append(merr)
            if mode is None:
                mode = "offline"

            status, serr = parse_enum(cell("status"), STATUS_ALIASES, "状态")
            if serr:
                errors.append(serr)
            if status is None:
                status = "scheduled"

            rate, rerr = parse_money(cell("rate"))
            if rerr:
                errors.append(rerr)

            row = {
                "row_no": r,
                "date": d,
                "time": t,
                "start_at": f"{d}T{t}" if (d and t) else None,
                "student": student,
                "duration_min": dur,
                "topic": _s(cell("topic")),
                "mode": mode,
                "location": _s(cell("location")),
                "rate": rate,
                "status": status,
                "errors": errors,
                "warnings": warnings,
            }
            rows.append(row)
            if len(rows) >= MAX_ROWS:
                notes.append(f"这份表超过 {MAX_ROWS} 行，后面的没有再读（一次别导太多，容易看不过来）")
                break

        if not rows:
            result["error"] = "表里没有可读的数据行（只有表头？）"
            return result

        result["ok"] = True
        result["rows"] = rows
        result["counts"] = {
            "total": len(rows),
            "ok": sum(1 for x in rows if not x["errors"]),
            "problem": sum(1 for x in rows if x["errors"]),
        }
        if filled_year:
            notes.append(f"有 {filled_year} 行没写年份，按 {default_year} 年处理")
        if from_range:
            notes.append(f"有 {from_range} 行的时长是从「19:00-20:00」这类区间里算出来的")
        if result["ignored_columns"]:
            notes.append("忽略了这几列：" + "、".join(result["ignored_columns"]))
        result["notes"] = notes

        # 排序：**没问题的排前面按时间排，有问题的沉到后面**（老师核对时先看能导的，
        # 有错的集中在一起也方便对着 Excel 改）。同级再按原行号，保证顺序稳定可预期。
        result["rows"].sort(key=lambda x: (1 if x["errors"] else 0,
                                           x["start_at"] or "9999",
                                           x["row_no"]))
        return result
    finally:
        try:
            wb.close()
        except Exception:
            pass


# ---------------------------------------------------------------- 模板
TEMPLATE_HEADERS = [label for _f, label, _r, _a, _h in COLUMNS]

TEMPLATE_SAMPLES = [
    ["2026-10-01", "19:00-20:30", "示例：张三", "", "三角函数图像变换", "线下", "教室 A", "", "已排课"],
    ["2026-10-03", "10:00", "示例：李四", 90, "导数与单调性", "线上", "", 400, "已排课"],
]

# (A 列, B 列)。**A 列非空时加粗**，用来当小标题；一行说明就只写 B 列。
TEMPLATE_NOTES = [
    ("这份表怎么填", ""),
    ("", "只有「日期」「开始时间」「学生」三列是必填的，其余留空也能导。"),
    ("", "列的顺序可以随意调换、可以自己加「备注」列（会被忽略），但表头名字要对得上。"),
    ("", "表头行上面可以有标题行（比如「10 月课表」），程序会自己往下找表头。"),
    ("", "只认 .xlsx。老的 .xls 请先在 Excel 里「另存为」成 .xlsx 再导入。"),
    ("", ""),
    ("每一列", ""),
    *[(f"「{label}」{'（必填）' if req else '（可空）'}", hint) for _f, label, req, _a, hint in COLUMNS],
    ("", ""),
    ("几种写法都认", ""),
    ("", "日期：2026-10-01 ／ 2026/10/1 ／ 2026年10月1日 ／ Excel 日期格式 ／ 10-01（补当年）"),
    ("", "时间：19:00 ／ 19：00 ／ 1900 ／ 19:00-20:30（会自动算出时长）"),
    ("", "时长：90 ／ 90分钟 ／ 1.5小时 ／ 1小时30分"),
    ("", "形式：线下 / 线上（也可以写面授、网课）"),
    ("", "状态：已排课 / 已完成 / 补课 / 请假 / 已取消 / 已调课"),
    ("", "单价：300 ／ ¥300 ／ 300元 ／ 300/小时"),
    ("", ""),
    ("导入时会发生什么", ""),
    ("", "1. 先「预览」：逐行告诉你哪几行能导、哪几行的学生系统里没有、哪几行和已排的课撞了。"),
    ("", "2. 你确认无误再点导入。学生名对不上的行，可以勾选「顺手建这几个学生」。"),
    ("", "3. 单价留空时，按学生档案里的默认单价快照到这节课上（以后改学生单价不影响历史）。"),
]


def build_template() -> bytes:
    """生成导入模板（.xlsx 字节）。两个工作表：「课表」可直接填，「填写说明」是这份说明。"""
    if not _HAS_OPENPYXL:
        raise RuntimeError("服务端没装 openpyxl，生成不了模板")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "课表"

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2563EB")
    note_font = Font(bold=True)
    sample_fill = PatternFill("solid", fgColor="F2F4F7")

    ws.append(TEMPLATE_HEADERS)
    for c in range(1, len(TEMPLATE_HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in TEMPLATE_SAMPLES:
        ws.append(row)
    # 示例行给个浅底色 + 说明「删掉这几行」
    for r in range(2, 2 + len(TEMPLATE_SAMPLES)):
        for c in range(1, len(TEMPLATE_HEADERS) + 1):
            ws.cell(row=r, column=c).fill = sample_fill
    ws.append([])
    tip_row = 2 + len(TEMPLATE_SAMPLES) + 1
    ws.cell(row=tip_row, column=1, value="↑ 上面两行是示例，填之前请删掉。本表从第几行开始填都行，表头别删。")
    ws.cell(row=tip_row, column=1).font = Font(italic=True, color="98A2B3")

    widths = [14, 14, 16, 14, 26, 12, 16, 12, 12]
    for i, w in enumerate(widths[: len(TEMPLATE_HEADERS)], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("填写说明")
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 78
    for col_a, col_b in TEMPLATE_NOTES:
        ws2.append([col_a, col_b])
        r = ws2.max_row
        if col_a:
            ws2.cell(row=r, column=1).font = note_font
        if col_b:
            ws2.cell(row=r, column=2).alignment = Alignment(wrap_text=False, vertical="center")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
